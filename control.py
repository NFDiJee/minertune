#!/usr/bin/env python3
"""Steuerzentrale: Webserver + Sweep-Steuerung in einem Programm.

- Sweep-Logik, Messkern und alle Sicherungen kommen unveraendert aus sweep_full.py (importiert).
- Im Ruhezustand KEIN Schreibzugriff auf den Miner; lesende Endpunkte brauchen kein PIN.
- Schreibende Endpunkte verlangen den Header X-Control-Pin (aus config.json).
- Ausnahme: POST /emergency-stop (Not-Aus) ohne PIN.
- Laeuft ein Sweep und wird er gestoppt / bricht ab / Prozess endet: Miner IMMER auf den
  Betriebspunkt vom Sweep-Start zurueck (sweep_full.restore im finally des Sweep-Threads).
Server: nur Standardbibliothek. openpyxl/reportlab (venv) nur fuer XLSX/PDF-Export.
"""

import argparse
import concurrent.futures
import copy
import csv
import glob
import fcntl
import hmac
import io
import ipaddress
import json
import os
import re
import signal
import socket
import struct
import sys
import threading
import time
import traceback
import urllib.error
import urllib.parse
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import tuner
import sweep as base
import sweep_full as sf

HERE = os.path.dirname(os.path.abspath(__file__))
RUNS_DIR = os.path.join(HERE, "runs")
DEFAULT_PIN = "CHANGEME"
DEFAULT_CONFIG = {
    "miner_url": "http://192.168.1.100/api/v1", "bind_host": "0.0.0.0", "bind_port": 8477,
    "control_pin": DEFAULT_PIN, "window_min_s": 600, "target_hits": 600, "window_max_s": 900,
    "freq_low_pct": 0.52, "freq_high_pct": 0.0, "allow_above_stock": False, "freq_step_factor": 25,
    "mv_step_factor": 5, "floor_pct": 0.75, "error_max_pct": 2.0, "hashrate_min_frac": 0.90,
    "soft_asic_c": 70, "soft_vr_c": 85, "hard_asic_c": 80, "hard_vr_c": 95, "input_v_min_frac": 0.95,
    "freq_verify_tol_mhz": 2,
    "settle_ma_tol_frac": 0.05, "warmup_max_s": 90, "language": "en",
}
PLAN_KEYS = ["freq_low_pct", "freq_high_pct", "allow_above_stock", "freq_step_factor", "mv_step_factor", "floor_pct"]
MAX_BODY = 64 * 1024

CFG = {}


class MinerContext:
    """Alles, was pro Miner existiert: Ziel-URL, verbundenes Profil, Sweep-Thread, Live-Zustand,
    geplante Matrix, Safe-Point, Sicherungs-Schwellen, letzter Lauf.

    Vorerst gibt es genau EINEN Kontext ("default"). Die Sweep-Engine (sweep_full) haelt ihren
    Laufzustand (LIVE/RUN/STATE/STOP/D) noch auf Modulebene; der Kontext greift ausschliesslich ueber
    self.engine darauf zu. Fuer mehrere Miner wird die Engine spaeter instanzbasiert (eine je Kontext)."""

    def __init__(self, cid, engine):
        self.id = cid
        self.engine = engine
        self.lock = threading.RLock()
        self.thread = None
        self.emergency = False
        self.emergency_at = None
        self.profile = None
        self.derived = None
        self.plan = None            # Kurzinfo des letzten /plan
        self.matrix = None          # geplante Matrix (aus /plan oder Sweep-Start) inkl. Laufbezug
        self.connected_at = None
        self.run_info = None
        self.events = []

    # --- Sicht auf den Engine-Zustand dieses Kontexts ---
    @property
    def miner_url(self):
        return self.engine.MINER_BASE

    @property
    def live(self):
        return self.engine.LIVE

    @property
    def run(self):
        return self.engine.RUN

    @property
    def stop_event(self):
        return self.engine.STOP

    @property
    def engine_state(self):
        return self.engine.STATE

    @property
    def safe_point(self):
        return self.engine.D.get("safe")

    @property
    def thresholds(self):
        e = self.engine
        return {"soft_asic_c": e.SOFT_ASIC_C, "soft_vr_c": e.SOFT_VR_C, "hard_asic_c": e.HARD_ASIC_C,
                "hard_vr_c": e.HARD_VR_C, "input_v_min": e.D.get("input_v_min")}

    @property
    def last_run(self):
        return {k: self.run.get(k) for k in ("timestamp", "finished_at", "status", "restored", "run_tag")}

    def sweep_running(self):
        t = self.thread
        return t is not None and t.is_alive()

    def add_event(self, line):
        with self.lock:
            self.events = (self.events + [line])[-30:]

    def public(self):
        """Kontrollblock fuer /state.json (ohne Lock/Thread-Objekte)."""
        with self.lock:
            return {"context_id": self.id, "emergency": self.emergency, "emergency_at": self.emergency_at,
                    "profile": copy.deepcopy(self.profile), "plan": copy.deepcopy(self.plan),
                    "connected_at": self.connected_at, "run_info": copy.deepcopy(self.run_info),
                    "events": list(self.events), "sweep_running": self.sweep_running(),
                    "miner_url": self.miner_url, "safe_point": self.safe_point, "thresholds": self.thresholds,
                    "last_run": self.last_run}

    # --- Matrix mit Fortschritt ---
    def set_matrix(self, pl, selected=None, source="plan", run_ts=None):
        freqs = sorted({q["freq"] for q in pl["points"]})
        if selected:
            freqs = [f for f in freqs if f in set(selected)]
        if pl.get("order") in ("desc", "absteigend"):     # "absteigend": Altlaeufe
            freqs = sorted(freqs, reverse=True)
        pts = {q["freq"]: q for q in pl["points"]}
        est = (pl.get("estimate") or {})
        per_f = None
        if est and pl["points"]:
            fast = est.get("fast") or {}
            per_f = (fast.get("duration_realistic_s") or est.get("duration_realistic_s") or 0) / max(len(pl["points"]), 1)
        with self.lock:
            self.matrix = {"source": source, "run_ts": run_ts, "sweep_mode": pl["sweep_mode"], "order": pl.get("order"),
                           "freqs": freqs, "plan_points": {str(f): pts[f] for f in freqs if f in pts},
                           "per_freq_est_s": per_f, "created_at": time.strftime("%Y-%m-%dT%H:%M:%S")}

    def matrix_status(self):
        with self.lock:
            m = copy.deepcopy(self.matrix)
        if not m:
            return None
        run, live = self.run, self.live
        running = self.sweep_running()
        belongs = m["source"] == "sweep" and m["run_ts"] and run.get("timestamp") == m["run_ts"]
        points = [q for q in (run.get("points") or []) if belongs and not str(q.get("label", "")).startswith("cal-")]
        vmin = (run.get("vmin") or {}) if belongs else {}
        cur = live.get("current") or {}
        cur_label = str((live.get("results") or [{}])[-1].get("label", "")) if live.get("results") else ""
        calibrating = belongs and running and not points and cur.get("phase") not in ("idle",) and \
            (run.get("gh_per_mhz") is None)
        cur_f = cur.get("frequency_mhz") if (belongs and running and not calibrating) else None
        phase = cur.get("phase")
        if phase == "warmup" and (cur.get("elapsed_s") or 0) < self.engine.SETTLE_PRECHECK_S and self.engine.GH_PER_MHZ:
            phase = "precheck"
        rows, done, active_idx, spent_done = [], 0, None, 0.0
        for i, f in enumerate(m["freqs"]):
            fp = [q for q in points if q.get("freq") == f]
            vm = vmin.get(str(f)) or {}
            sub = []
            for q in fp:
                tag = q.get("tag") or legacy_tag(str(q.get("label", "")))
                st = "gueltig" if q.get("valid") else ("vorab_ok" if q.get("status") == "precheck_pass" else "verworfen")
                sub.append({"mv": q.get("mv"), "tag": tag, "status": st, "stage": q.get("stage"),
                            "frac": q.get("frac_of_expected"), "ths": q.get("hashrate_local_ths"),
                            "jth": q.get("jth_local"), "wall": q.get("wall_avg"), "reason": q.get("reason"),
                            "vmin": vm.get("vmin") is not None and q.get("mv") == vm.get("vmin") and bool(q.get("valid"))})
            spent = sum((q.get("point_s") or 0) + 5 for q in fp)
            if cur_f == f and phase not in ("idle", "done"):
                status = "aktiv"
                active_idx = i
                sub.append({"mv": cur.get("core_mv"), "tag": "", "status": "aktiv", "phase": phase,
                            "elapsed_s": cur.get("elapsed_s"), "remaining_s": cur.get("remaining_s")})
            elif vm.get("vmin") is not None:
                status = "fertig"
            elif fp or vm:
                status = "verworfen"
            else:
                status = "uebersprungen" if (belongs and not running) else "wartet"
            if status in ("fertig", "verworfen"):
                done += 1
                spent_done += spent
            rec = next((q for q in fp if q.get("valid") and q.get("mv") == vm.get("vmin")), None)
            pp = m["plan_points"].get(str(f)) or {}
            rows.append({"freq": f, "status": status, "sub": sub, "vmin": vm.get("vmin"), "found": vm.get("found"),
                         "start_mv": vm.get("start_mv") or (fp[0].get("mv") if fp else None) or
                                     (cur.get("core_mv") if cur_f == f else None) or
                                     (pp.get("mv_start") if m["sweep_mode"] not in ("efficiency", "target_hashrate") else None),
                         "mv_floor": pp.get("mv_floor"), "mv_top": pp.get("mv_start"),
                         "direction": "auf" if m["sweep_mode"] in ("efficiency", "target_hashrate") else "ab",
                         "role": pp.get("role"), "ths": rec.get("hashrate_local_ths") if rec else None,
                         "jth": rec.get("jth_local") if rec else None, "spent_s": round(spent)})
        n = len(m["freqs"])
        x = done + (1 if active_idx is not None else 0)
        avg = spent_done / done if done else m.get("per_freq_est_s")
        remaining = None
        if avg and running and belongs:
            act_spent = rows[active_idx]["spent_s"] + (cur.get("elapsed_s") or 0) if active_idx is not None else 0
            remaining = max(0.0, avg * (n - done) - act_spent)
        return {"source": m["source"], "sweep_mode": m["sweep_mode"], "order": m["order"], "rows": rows,
                "progress": {"x": x, "n": n, "done": done, "calibrating": bool(calibrating),
                             "text": ("Calibration" if calibrating else f"Frequency {x} of {n}"),
                             "remaining_s": round(remaining) if remaining is not None else None},
                "running": running and bool(belongs), "early_stop": run.get("early_stop") if belongs else None}


contexts = {"default": MinerContext("default", sf)}


def ctx(cid="default"):
    return contexts[cid]


def clog(msg):
    line = f"{time.strftime('%H:%M:%S')}  {msg}"
    print(line, flush=True)
    ctx().add_event(line)


# ----------------------------------------------------------------------------
# Konfiguration
# ----------------------------------------------------------------------------
def load_config(path):
    created = False
    if not os.path.exists(path):
        with open(path, "w", encoding="utf-8") as f:
            json.dump(DEFAULT_CONFIG, f, indent=2)
        created = True
    with open(path, encoding="utf-8") as f:
        cfg = json.load(f)
    if "language" not in cfg:            # Feld nachruesten (atomar, uebrige Felder + Rechte unveraendert)
        cfg["language"] = DEFAULT_CONFIG["language"]
        tmp = path + ".tmp"
        mode = os.stat(path).st_mode & 0o777
        fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, mode)
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=2)
            f.flush()
            os.fsync(f.fileno())
        os.replace(tmp, path)
    merged = dict(DEFAULT_CONFIG, **cfg)
    return merged, created


# ----------------------------------------------------------------------------
# UI-Sprachen: lang/<code>.json (automatisch gefunden, Anzeigename aus "_name")
# ----------------------------------------------------------------------------
LANG_DIR = os.path.join(HERE, "lang")
LANG_CODE_RE = re.compile(r"^[a-z]{2,3}(-[A-Za-z0-9]{2,8})?$")


def lang_list():
    out = []
    for path in sorted(glob.glob(os.path.join(LANG_DIR, "*.json"))):
        code = os.path.splitext(os.path.basename(path))[0]
        if not LANG_CODE_RE.match(code):
            continue
        try:
            with open(path, encoding="utf-8") as f:
                name = json.load(f).get("_name") or code
        except Exception:
            continue
        out.append({"code": code, "name": name})
    return out


def lang_table(code):
    if not LANG_CODE_RE.match(code or ""):
        return None
    path = os.path.join(LANG_DIR, code + ".json")
    if not os.path.isfile(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


LEGACY_TAGS = (("Knie-Reserve", "knee_reserve"), ("Reserve", "reserve"), ("nachtasten", "probe_down"),
               ("Vorabcheck", "precheck"))
LEGACY_WALLSRC = {"gemessen": "measured", "geschaetzt": "estimated", "gemischt": "mixed"}


def legacy_tag(label):
    """Punkt-Tag aus dem Label alter Laeufe (deutsche Label-Zusaetze) -> sprachneutraler Code."""
    return next((code for word, code in LEGACY_TAGS if f"({word})" in label), "")


def exc_body(e, fallback_prefix=False):
    """Fehlerantwort aus einer Exception: Engine-Fehler mit code/params -> uebersetzbar, sonst Klartext."""
    code, params = getattr(e, "code", None), getattr(e, "params", None)
    if isinstance(code, str) and isinstance(params, dict):
        return err(code, str(e), **params)
    return {"error": f"{type(e).__name__}: {e}" if fallback_prefix else str(e)}


def err(code, text, **params):
    """API-Fehler: englischer Text + Code (+ Parameter), den das Frontend uebersetzt."""
    return {"error": text, "code": code, "params": params}


def apply_config():
    sf.configure(CFG["miner_url"], {k: v for k, v in CFG.items() if k in sf.CONFIG_KEYS})


CONFIG_PATH = {"path": os.path.join(HERE, "config.json")}
PIN_LOCK = threading.Lock()
PIN_MIN_LEN, PIN_MAX_LEN = 4, 128


def pin_set():
    """PIN gilt als eingerichtet, wenn control_pin vorhanden, nicht leer und nicht 'CHANGEME' ist."""
    v = CFG.get("control_pin")
    return isinstance(v, str) and v.strip() != "" and v != DEFAULT_PIN


def pin_default():
    return not pin_set()


def validate_new_pin(new_pin):
    if not isinstance(new_pin, str) or new_pin.strip() == "":
        return "new_pin must not be empty"
    if new_pin == DEFAULT_PIN:
        return "new_pin must not be 'CHANGEME'"
    if len(new_pin) < PIN_MIN_LEN:
        return f"new_pin must have at least {PIN_MIN_LEN} characters"
    if len(new_pin) > PIN_MAX_LEN:
        return f"new_pin must have at most {PIN_MAX_LEN} characters"
    return None


def write_pin_to_config(new_pin, path=None, update_runtime=True):
    """Nur control_pin in config.json ersetzen (uebrige Felder unveraendert), atomar via Temp-Datei + rename,
    danach Laufzeit-Konfig aktualisieren -> neuer PIN gilt sofort, ohne Neustart."""
    path = path or CONFIG_PATH["path"]
    with open(path, encoding="utf-8") as f:
        on_disk = json.load(f)
    on_disk["control_pin"] = new_pin
    tmp = path + ".tmp"
    fd = os.open(tmp, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)   # PIN-Datei nur fuer den Besitzer lesbar
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        json.dump(on_disk, f, indent=2)
        f.flush()
        os.fsync(f.fileno())
    os.replace(tmp, path)
    os.chmod(path, 0o600)
    if update_runtime:
        CFG["control_pin"] = new_pin


def reload_pin_from_file():
    """Laufzeit-Reload nur des PINs aus config.json (SIGHUP) - kein Neustart noetig."""
    with open(CONFIG_PATH["path"], encoding="utf-8") as f:
        CFG["control_pin"] = json.load(f).get("control_pin", DEFAULT_PIN)
    clog("Control PIN reloaded from config.json (SIGHUP): " + ("set" if pin_set() else "NOT set"))


# ----------------------------------------------------------------------------
# Miner lesen / Plan-Vorschau (nur GET)
# ----------------------------------------------------------------------------
def read_status_once():
    st = tuner.get_status()
    w, src = sf.effective_wall_power(st)
    st["_effective_wall_power_w"], st["_wall_power_source"] = w, src
    return st


def profile_from(st):
    return {k: st.get(k) for k in ["profile", "profile_id", "stock_frequency_mhz", "stock_core_mv", "frequency_min",
                                   "frequency_max", "frequency_step", "core_min", "core_max", "core_step",
                                   "local_hashrate_difficulty", "input_voltage_v", "current_frequency_mhz", "core_mv",
                                   "local_hashrate_ghs", "job_interval_ms", "mining_enabled", "asic_temp_c",
                                   "vr_temp_c", "wall_power_measured"]}


def preview(p, d):
    gh_est = (p["local_hashrate_ghs"] / p["current_frequency_mhz"]
              if isinstance(p.get("local_hashrate_ghs"), (int, float)) and p.get("current_frequency_mhz") else None)
    diff = p.get("local_hashrate_difficulty")
    n = len(d["mv_steps"])
    n_real = min(n, sf.REALISTIC_STEPS)
    rows, worst, real = [], 0.0, 0.0
    for f in d["freqs"]:
        win, _ = sf.window_estimate(f, gh_est, diff)
        per = sf.WARMUP_EST_S + win + sf.OVERHEAD_EST_S
        worst += n * per
        real += n_real * per
        rows.append({"frequency_mhz": f, "pct_of_stock": round(f / p["stock_frequency_mhz"] * 100, 1),
                     "mv_start": d["mv_start"], "mv_floor": d["mv_floor"], "steps": n,
                     "expected_ths": round(f * gh_est / 1000, 3) if gh_est else None, "window_s": round(win)})
    cal = 2 * (sf.WARMUP_EST_S + sf.CAL_WINDOW_S + sf.OVERHEAD_EST_S)
    return {"rows": rows, "points_worst_case": n * len(rows), "points_realistic": n_real * len(rows),
            "duration_worst_s": round(worst + cal), "duration_realistic_s": round(real + cal),
            "derived": {k: v for k, v in d.items() if k not in ("formula",)}, "formula": d["formula"],
            "params": {k: getattr(sf, sf.CONFIG_KEYS[k]) for k in PLAN_KEYS}}


def connect_and_derive():
    st = read_status_once()
    p = profile_from(st)
    missing = [k for k in ["stock_frequency_mhz", "stock_core_mv", "frequency_min", "frequency_max",
                           "frequency_step", "core_min", "core_max", "core_step"] if not isinstance(p[k], (int, float))]
    if missing:
        raise RuntimeError(f"Profile incomplete: {missing}")
    d = sf.derive(p)
    c = ctx()
    with c.lock:
        c.profile, c.derived, c.connected_at = p, d, time.strftime("%Y-%m-%dT%H:%M:%S")
    return st, p, d


def sweep_running():
    return ctx().sweep_running()


# ----------------------------------------------------------------------------
# Sweep-Thread (Engine aus sweep_full)
# ----------------------------------------------------------------------------
RUN_OVERRIDE_LIMITS = {"window_min_s": (30, 3600), "window_max_s": (30, 3600), "target_hits": (10, 5000)}
RUN_OVERRIDE_ATTR = {"window_min_s": "WINDOW_MIN_S", "window_max_s": "WINDOW_MAX_S", "target_hits": "TARGET_HITS"}


def sweep_worker(req, cid="default"):
    """Gerichtete Suche (Engine des Kontexts) im Hintergrund. req: validierte Start-Parameter."""
    c = ctx(cid)
    sf = c.engine      # Engine dieses Kontexts (vorerst das Modul sweep_full)
    p = None
    try:
        sf.reset_run()
        apply_config()                                   # Konfig-Defaults ...
        for k, v in req["run_overrides"].items():        # ... nur fuer DIESEN Lauf ueberschreiben
            setattr(sf, RUN_OVERRIDE_ATTR[k], v)
        if req.get("run_tag"):
            sf.RUN_PATH = os.path.join(RUNS_DIR, f"{req['run_tag']}_{time.strftime('%Y%m%d_%H%M%S')}.json")
        sf.RUN.update(timestamp=time.strftime("%Y-%m-%dT%H:%M:%S"), started_by="control.py",
                      selected_freqs=req["selected"], run_tag=req.get("run_tag"), note=req.get("note"),
                      run_overrides=req["run_overrides"], request={k: v for k, v in req.items() if k != "note"})
        st, p = sf.read_profile()                        # Startpunkt lesen -> Ruecksetzziel
        sf.D.update(sf.derive(p))
        pl = sf.plan_matrix(p, req["sweep_mode"], req["resolution"], req.get("target_ths"),
                            req.get("allow_above_stock"), req.get("freq_high_pct"), None, None, req.get("plan_overrides"))
        if pl["requires_gate"] and not req.get("confirm_gate"):
            raise RuntimeError("Plan leaves the factory range - start only with confirm_gate=true")
        pl["estimate"] = sf.plan_estimate(pl, p.get("local_hashrate_difficulty"))
        c.set_matrix(pl, req["selected"], source="sweep", run_ts=sf.RUN["timestamp"])
        clog(f"Sweep started: mode {req['sweep_mode']}, {req['resolution']}, full_measure={req['full_measure']}, "
             f"selection {req['selected'] or 'all'}, start point {sf.D['safe'][0]}/{sf.D['safe'][1]}"
             + (f", run overrides {req['run_overrides']}" if req["run_overrides"] else "")
             + (f", tag {req['run_tag']}" if req.get("run_tag") else ""))
        sf.run_plan(st, p, sf.D, pl, req["selected"], req["full_measure"])
    except sf.Stopped:
        sf.RUN["status"] = "emergency_stop" if c.emergency else "stopped"
        sf.log("Sweep stopped (web control).")
    except SystemExit:  # sweep_full.emergency_stop (HARD-Temp) -> Status bereits gesetzt
        c.emergency = True
        c.emergency_at = time.strftime("%Y-%m-%dT%H:%M:%S")
    except base.Deadman as e:
        sf.log(f"ABORT: {e}")
        sf.RUN["status"] = "deadman"
    except Exception as e:
        sf.log(f"ERROR: {type(e).__name__}: {e}")
        sf.RUN["status"] = f"error: {e}"
        traceback.print_exc()
    finally:
        try:
            if p is not None:
                if c.emergency:
                    sf.STATE["emergency"] = True
                sf.restore(p)
        finally:
            sf.RUN["finished_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
            try:
                sf.save_run()
                sf.log(f"Run JSON saved: {sf.RUN_PATH}")
            except Exception as e:
                sf.log(f"!!! Run JSON could not be saved: {e}")
            sf.publish_live(idx=None, phase="idle")
            clog(f"Sweep finished: status {sf.RUN.get('status')}, restored={sf.RUN.get('restored')}")


# ----------------------------------------------------------------------------
# Einzelpunkt setzen (PIN; save=true nur mit confirm_save)
# ----------------------------------------------------------------------------
def set_single_point(freq, mv, save):
    st = read_status_once()
    p = profile_from(st)
    sf.check_wall(freq, mv, p)
    sfq, smv = p["stock_frequency_mhz"], p["stock_core_mv"]
    df, dm = abs(freq - sfq) / sfq, abs(mv - smv) / smv
    payload = {"frequency_mhz": freq, "core_mv": mv, "save": bool(save)}
    if df > sf.DANGER_PCT or dm > sf.DANGER_PCT:
        payload["danger_acknowledged"] = True
    code, text = base.post_json(sf.TUNING, payload)
    time.sleep(1.5)
    st2 = read_status_once()
    ok, _, meas, vtext = sf.verify_state(st2, freq, mv, p)
    clog(f"Single point set: {json.dumps(payload)} -> HTTP {code}; {vtext}")
    return {"payload": payload, "http": code, "response": text.strip(), "verify_ok": ok,
            "current_frequency_mhz": st2.get("current_frequency_mhz"), "measured_core_mv": meas,
            "danger": {"freq_pct": round(df * 100, 1), "mv_pct": round(dm * 100, 1)}}


# ----------------------------------------------------------------------------
# Miner-URL, Subnetze, Suche, Identify
# ----------------------------------------------------------------------------
SCAN_TIMEOUT_S = 1.0
SCAN_WORKERS = 64
SEARCH_LOCK = threading.Lock()
HOST_RE = re.compile(r"^[A-Za-z0-9.\-]+$")


def normalize_miner_url(eingabe):
    """'192.168.1.100' | '192.168.1.100:8080' | 'http://host[/api/v1]' -> 'http://<host>[:port]/api/v1'.
    Vorhandenes http/https wird respektiert, Standardports (80/443) entfallen, /api/v1 wird angehaengt."""
    raw = (eingabe or "").strip()
    if not raw:
        raise ValueError("empty input")
    if not re.match(r"^[A-Za-z][A-Za-z0-9+.-]*://", raw):
        raw = "http://" + raw
    u = urllib.parse.urlsplit(raw)
    if u.scheme not in ("http", "https"):
        raise ValueError(f"scheme {u.scheme!r} not allowed (http/https only)")
    host = u.hostname
    if not host or not HOST_RE.match(host):
        raise ValueError(f"invalid host in {eingabe!r}")
    try:
        port = u.port
    except ValueError:
        raise ValueError(f"invalid port in {eingabe!r}")
    netloc = host if port is None or (u.scheme, port) in (("http", 80), ("https", 443)) else f"{host}:{port}"
    path = (u.path or "").rstrip("/")
    if not path.endswith("/api/v1"):
        path = path + "/api/v1"
    return f"{u.scheme}://{netloc}{path}"


def _default_route_iface():
    try:
        with open("/proc/net/route") as f:
            for line in f.readlines()[1:]:
                parts = line.split()
                if len(parts) > 2 and parts[1] == "00000000" and int(parts[3], 16) & 2:
                    return parts[0]
    except OSError:
        pass
    return None


def _iface_names():
    """Interface-Namen. socket.if_nameindex() braucht in glibc einen AF_NETLINK-Socket; der ist unter
    systemd RestrictAddressFamilies=AF_INET AF_INET6 (minertune.service) verboten -> OSError Errno 97.
    Dann ohne Socket aus /proc/net/dev lesen. Nichts lesbar -> leere Liste statt Absturz."""
    try:
        return [name for _, name in socket.if_nameindex()]
    except OSError:
        pass
    try:
        with open("/proc/net/dev") as f:
            return [ln.split(":", 1)[0].strip() for ln in f.readlines()[2:] if ":" in ln]
    except OSError:
        return []


def _iface_ipv4(name):
    with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:     # bewusst nur IPv4 (/24-Scan)
        req = struct.pack("256s", name.encode()[:15])
        addr = socket.inet_ntoa(fcntl.ioctl(s.fileno(), 0x8915, req)[20:24])     # SIOCGIFADDR
        mask = socket.inet_ntoa(fcntl.ioctl(s.fileno(), 0x891B, req)[20:24])     # SIOCGIFNETMASK
    return addr, mask


def local_subnets():
    """Eigene IPv4-Interfaces -> /24-Netze mit Label; Default-Route-Interface = empfohlen."""
    default_if = _default_route_iface()
    out, seen = [], set()
    for name in _iface_names():
        try:
            addr, mask = _iface_ipv4(name)          # Interface ohne IPv4 / nicht abfragbar -> ueberspringen
            ip = ipaddress.IPv4Address(addr)
        except (OSError, ValueError):
            continue
        if ip.is_loopback or ip.is_link_local:
            continue
        net = ipaddress.IPv4Network(f"{addr}/24", strict=False)
        if str(net) in seen:
            continue
        seen.add(str(net))
        if name == default_if:
            label = "aktiv/LAN"
        elif ip in ipaddress.IPv4Network("172.16.0.0/12"):
            label = "docker"
        elif ip in ipaddress.IPv4Network("100.64.0.0/10"):
            label = "vpn/tailscale"
        else:
            label = "sonstiges"
        out.append({"subnet": str(net), "interface": name, "address": addr, "netmask": mask,
                    "label": label, "recommended": name == default_if})
    out.sort(key=lambda x: (not x["recommended"], x["label"] != "sonstiges", x["subnet"]))
    return out


def probe_miner(ip):
    """Einzelnes GET auf http://<ip>/api/v1/status (nur lesend). -> Trefferdict oder None.
    ip ist ein IPv4-Literal -> Verbindung immer ueber AF_INET. Jeder Fehler (auch OSError/Errno 97)
    markiert nur diesen Host als nicht erreichbar; der Scan laeuft weiter."""
    try:
        req = urllib.request.Request(f"http://{ip}/api/v1/status", method="GET", headers={"Accept": "application/json"})
        with urllib.request.urlopen(req, timeout=SCAN_TIMEOUT_S) as r:
            if r.status != 200:
                return None
            st = json.loads(r.read(256 * 1024).decode("utf-8", errors="replace"))
    except Exception:
        return None
    if not isinstance(st, dict) or "profile" not in st or "stock_frequency_mhz" not in st:
        return None
    h = st.get("local_hashrate_ghs")
    return {"ip": ip, "profile": st.get("profile"), "profile_id": st.get("profile_id"),
            "current_frequency_mhz": st.get("current_frequency_mhz"), "core_mv": st.get("core_mv"),
            "local_hashrate_ths": round(h / 1000, 3) if isinstance(h, (int, float)) else None,
            "mining_enabled": st.get("mining_enabled"), "asic_temp_c": st.get("asic_temp_c"),
            "vr_temp_c": st.get("vr_temp_c")}


PRIVATE_NETS = [ipaddress.IPv4Network(n) for n in ("10.0.0.0/8", "172.16.0.0/12", "192.168.0.0/16")]


class SubnetError(ValueError):
    """Ungueltige/nicht erlaubte Subnetz-Eingabe: code + params fuer die Uebersetzung (err.<code>)."""

    def __init__(self, code, text, **params):
        self.code, self.params = code, params
        super().__init__(text)


def parse_custom_subnet(text):
    """Freie Eingabe -> IPv4Network /24. Erlaubt: 'x.y.z.0/24', 'x.y.z' (3 Oktette), 'x.y.z.w' (dessen /24).
    Nur private Bereiche (RFC 1918) und hoechstens 254 Hosts - kein Scannen fremder/oeffentlicher Netze."""
    raw = str(text or "").strip()
    if re.fullmatch(r"\d{1,3}\.\d{1,3}\.\d{1,3}\.?", raw):
        raw = raw.rstrip(".") + ".0/24"
    elif re.fullmatch(r"\d{1,3}(\.\d{1,3}){3}", raw):
        raw += "/24"
    try:
        net = ipaddress.IPv4Network(raw, strict=False)
    except ValueError:
        raise SubnetError("subnet_invalid", f"invalid subnet {text!r} (use x.y.z.0/24, private ranges only)", input=str(text))
    if net.prefixlen != 24:
        raise SubnetError("subnet_invalid", f"invalid subnet {text!r}: only /24 is allowed", input=str(text))
    if not any(net.subnet_of(p) for p in PRIVATE_NETS):
        raise SubnetError("subnet_public", f"{net} is not a private network - only private subnets can be scanned",
                          input=str(net))
    return net


def search_miners(which="auto"):
    subs = local_subnets()
    if which in (None, "", "auto"):
        nets = [x["subnet"] for x in subs if x["recommended"]] or [x["subnet"] for x in subs[:1]]
    elif which == "all":
        nets = [x["subnet"] for x in subs]
    elif str(which).strip() in {x["subnet"] for x in subs}:     # erkanntes eigenes Netz (Auswahlliste)
        nets = [str(which).strip()]
    else:                                                      # frei eingegeben: auch ohne eigenes Interface
        nets = [str(parse_custom_subnet(which))]
    if not nets:
        raise ValueError("no IPv4 network interface found for the miner search - enter the miner IP directly "
                         "or choose a subnet under 'Advanced'")
    hosts = [str(h) for n in nets for h in ipaddress.IPv4Network(n).hosts()]
    t0 = time.monotonic()
    clog(f"Miner search: {', '.join(nets)} ({len(hosts)} addresses, {SCAN_WORKERS} parallel, timeout {SCAN_TIMEOUT_S}s)")
    with concurrent.futures.ThreadPoolExecutor(max_workers=SCAN_WORKERS) as ex:
        hits = [r for r in ex.map(probe_miner, hosts) if r]
    hits.sort(key=lambda x: ipaddress.IPv4Address(x["ip"]))
    dur = time.monotonic() - t0
    clog(f"Miner search done: {len(hits)} hits in {dur:.1f}s")
    return {"subnets": nets, "scanned": len(hosts), "duration_s": round(dur, 2), "hits": hits}


def identify_miner(target):
    """Display 3x blinken lassen. Nutzt AUSSCHLIESSLICH /display (kein Tuning)."""
    url = normalize_miner_url(target)
    with urllib.request.urlopen(urllib.request.Request(url + "/status", method="GET"), timeout=10) as r:
        st = json.loads(r.read().decode("utf-8"))
    if not st.get("display_available"):
        return {"identified": False, "reason": "no display available", "code": "no_display", "miner_url": url}
    original = bool(st.get("display_enabled", True))
    display_url = url + "/display"
    try:
        for _ in range(3):
            base.post_json(display_url, {"enabled": False})
            time.sleep(0.5)
            base.post_json(display_url, {"enabled": True})
            time.sleep(0.5)
    finally:
        base.post_json(display_url, {"enabled": original})
    clog(f"Identify: {url} display blinked 3x, restored to enabled={original}")
    return {"identified": True, "miner_url": url, "display_restored_to": original}


# ----------------------------------------------------------------------------
# Runs / Export
# ----------------------------------------------------------------------------
RUN_ID_RE = re.compile(r"^[A-Za-z0-9_\-]{1,80}$")
EXPORT_COLS = ["label", "freq", "mv", "measured_mv", "valid", "reason", "hashrate_local_ths", "hashrate_valid_ths",
               "expected_ths", "frac_of_expected", "wall_avg", "wall_power_source", "rail_avg", "jth_local",
               "jth_valid", "error_pct", "asic_temp_max", "vr_temp_max", "settle_s", "dur", "dvalid"]


def run_files():
    files = [f for f in glob.glob(os.path.join(RUNS_DIR, "*.json"))
             if not os.path.basename(f).startswith("live_state")]
    return sorted(files, key=os.path.getmtime, reverse=True)


def load_run(run_id):
    if not RUN_ID_RE.match(run_id or ""):
        return None
    path = os.path.join(RUNS_DIR, run_id + ".json")
    if not os.path.isfile(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def run_summary(path):
    rid = os.path.splitext(os.path.basename(path))[0]
    try:
        with open(path, encoding="utf-8") as f:
            d = json.load(f)
    except Exception as e:
        return {"id": rid, "error": str(e)}
    pts = d.get("points", [])
    return {"id": rid, "run_tag": d.get("run_tag"), "note": d.get("note"), "sweep_mode": d["plan"].get("sweep_mode") if isinstance(d.get("plan"), dict) else None,
            "timestamp": d.get("timestamp"), "finished_at": d.get("finished_at"), "status": d.get("status"),
            "profile": d.get("profile"), "profile_id": d.get("profile_id"), "points": len(pts),
            "valid_points": sum(1 for q in pts if q.get("valid")), "gh_per_mhz": d.get("gh_per_mhz"),
            "sweet_spots": d.get("sweet_spots"), "restored": d.get("restored"),
            "mtime": time.strftime("%Y-%m-%dT%H:%M:%S", time.localtime(os.path.getmtime(path)))}


def _pt_summary(q, role=None):
    return {"freq": q.get("freq"), "mv": q.get("mv"), "hashrate_ths": q.get("hashrate_local_ths"),
            "wall_w": q.get("wall_avg"), "jth_local": q.get("jth_local"), "frac_of_expected": q.get("frac_of_expected"),
            "wall_power_quelle": q.get("wall_power_source") or "legacy", "role": role,
            "label": q.get("label")}


def result_points(d):
    """Je Frequenz die gueltigen Ergebnispunkte: Vmin (erster versorgter Punkt) und Reserve (Vmin + 1 Stufe).
    Verworfene Zwischenstufen und Kalibrierpunkte werden nicht betrachtet. Funktioniert fuer alle Laufformate."""
    pts = [q for q in d.get("points", []) if q.get("valid") and q.get("hashrate_local_ths") is not None
           and not str(q.get("label", "")).startswith("cal-")]
    vmin = d.get("vmin") or {}
    out = []
    for f in sorted({q["freq"] for q in pts}):
        valid = sorted((q for q in pts if q["freq"] == f), key=lambda q: q["mv"])
        by_mv = {q["mv"]: q for q in valid}
        v = vmin.get(str(f))
        v = v if isinstance(v, dict) else ({"vmin": v} if isinstance(v, (int, float)) else {})
        found = v.get("found", v.get("vmin"))
        if found not in by_mv:
            found = valid[0]["mv"]                       # Altlauf ohne Vmin-Angabe: niedrigster gueltiger Punkt
        res_mv = v.get("vmin") if v.get("vmin") not in (None, found) and v.get("vmin") in by_mv else \
            next((m for m in sorted(by_mv) if m > found), None)   # Reserve = naechste gueltige Stufe darueber
        vpt = by_mv[found]
        rpt = by_mv.get(res_mv) if res_mv is not None else None
        out.append({"freq": f, "vmin": vpt, "reserve": rpt, "recommended": rpt or vpt})
    return out


BEST_TIE_FRAC = 0.01   # Hashrates innerhalb 1 % gelten als gleichwertig -> niedrigster J/TH gewinnt


def best_for(d, run_id, target):
    rows = result_points(d)
    res = {"run_id": run_id, "target_ths": target, "candidates": len(rows), "chosen": None, "vmin_alternative": None,
           "below": None, "message": None, "caution": None,
           "rule": "smallest hashrate >= target (recommendation = reserve point Vmin+1 if measured); "
                   "among equivalent ones (< 1 %) the lowest J/TH"}
    tag = str(d.get("run_tag") or run_id)
    if "kurz" in tag:
        res["caution"] = ("Verification run with short measurement windows - values statistically uncertain, "
                          "do not use for permanent decisions.")
        res["caution_code"] = "short_run"
    if not rows:
        res["message"] = "Run contains no valid result points."
        res["message_code"] = "no_points"
        return res
    ok = [r for r in rows if r["recommended"]["hashrate_local_ths"] >= target]
    below = [r for r in rows if r["recommended"]["hashrate_local_ths"] < target]
    if below:
        b = max(below, key=lambda r: r["recommended"]["hashrate_local_ths"])
        res["below"] = _pt_summary(b["recommended"], "reserve" if b["reserve"] else "vmin")
    if not ok:
        top = max(rows, key=lambda r: r["recommended"]["hashrate_local_ths"])["recommended"]["hashrate_local_ths"]
        res["message"] = f"Target {target:g} TH/s is above the highest measured point ({top:.3f} TH/s)"
        res["message_code"] = "above_max"
        res["message_params"] = {"target": target, "top": top}
        return res
    hmin = min(r["recommended"]["hashrate_local_ths"] for r in ok)
    tied = [r for r in ok if r["recommended"]["hashrate_local_ths"] <= hmin * (1 + BEST_TIE_FRAC)]
    c = min(tied, key=lambda r: (r["recommended"].get("jth_local") or 1e9))
    res["chosen"] = _pt_summary(c["recommended"], "reserve_recommended" if c["reserve"] else "vmin_no_reserve")
    if c["reserve"]:
        res["vmin_alternative"] = _pt_summary(c["vmin"], "vmin_alternative")
    return res


def export_rows(d):
    rows = []
    for q in d.get("points", []):
        r = {k: q.get(k) for k in EXPORT_COLS}
        r["reason"] = sf.msg_text(q.get("reason"))       # Code-Objekte -> englischer Klartext; Altlauf-Text unveraendert
        r["wall_power_source"] = LEGACY_WALLSRC.get(r["wall_power_source"], r["wall_power_source"])
        if r["wall_power_source"] is None and q.get("wall_avg") is not None:
            r["wall_power_source"] = "unknown (legacy run)"
        rows.append(r)
    return rows


def export_csv(d):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=EXPORT_COLS)
    w.writeheader()
    for r in export_rows(d):
        w.writerow(r)
    return buf.getvalue().encode("utf-8")


def _ss_rows(d):
    out = []
    for key, name in [("efficiency", "Efficiency"), ("knee", "Performance knee"), ("compromise", "Compromise")]:
        s = (d.get("sweet_spots") or {}).get(key)
        if s:
            out.append([name, s.get("freq"), s.get("mv"), s.get("hashrate_local_ths"), s.get("wall_avg"), s.get("jth_local")])
    return out


def export_xlsx(d, run_id):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Points"
    ws.append(EXPORT_COLS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in export_rows(d):
        ws.append([r[k] for k in EXPORT_COLS])
    ws2 = wb.create_sheet("Sweet Spots")
    ws2.append(["Type", "MHz", "mV", "TH/s", "Wall W", "J/TH"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in _ss_rows(d):
        ws2.append(r)
    ws3 = wb.create_sheet("Info")
    for k in ["timestamp", "finished_at", "status", "miner_url", "profile", "profile_id", "gh_per_mhz",
              "job_interval_ms", "restored"]:
        ws3.append([k, str(d.get(k))])
    ws3.append(["run_id", run_id])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf(d, run_id):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=28, bottomMargin=28)
    ss = getSampleStyleSheet()
    f2 = lambda v, nd=2: f"{v:.{nd}f}" if isinstance(v, (int, float)) else ("" if v is None else str(v))
    story = [Paragraph(f"Miner sweep {run_id}", ss["Title"]),
             Paragraph(f"Profile: {d.get('profile')} ({d.get('profile_id')}) &nbsp; Start: {d.get('timestamp')} &nbsp; "
                       f"End: {d.get('finished_at')} &nbsp; Status: {d.get('status')} &nbsp; "
                       f"GH/MHz: {f2(d.get('gh_per_mhz'))}", ss["Normal"]), Spacer(1, 10)]
    ssr = _ss_rows(d)
    if ssr:
        t = Table([["Sweet Spot", "MHz", "mV", "TH/s", "Wall W", "J/TH"]] +
                  [[r[0], r[1], r[2], f2(r[3], 3), f2(r[4], 1), f2(r[5])] for r in ssr])
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3d2b")),
                               ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 9)]))
        story += [t, Spacer(1, 12)]
    cols = ["label", "freq", "mv", "measured_mv", "hashrate_local_ths", "wall_avg", "wall_power_source", "jth_local",
            "error_pct", "asic_temp_max", "vr_temp_max", "valid", "reason"]
    data = [["Point", "MHz", "mV", "meas. mV", "TH/s", "Wall W", "Source", "J/TH", "err %", "ASIC", "VR", "valid", "Reason"]]
    for r in export_rows(d):
        data.append([r["label"], r["freq"], r["mv"], f2(r["measured_mv"], 1), f2(r["hashrate_local_ths"], 3),
                     f2(r["wall_avg"], 1), r["wall_power_source"] or "", f2(r["jth_local"]), f2(r["error_pct"]),
                     f2(r["asic_temp_max"], 1), f2(r["vr_temp_max"], 1), "yes" if r["valid"] else "no",
                     (r["reason"] or "")[:60]])
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#263241")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 7)]))
    story.append(t)
    doc.build(story)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# HTTP
# ----------------------------------------------------------------------------
class Handler(BaseHTTPRequestHandler):
    server_version = "MinerControl/1"

    def log_message(self, fmt, *args):
        pass

    def _send(self, code, obj=None, ctype="application/json", body=None, filename=None):
        if body is None:
            body = json.dumps(obj, ensure_ascii=False, default=str).encode("utf-8")
        elif isinstance(body, str):
            body = body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        if filename:
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.end_headers()
        self.wfile.write(body)

    def _body(self):
        n = int(self.headers.get("Content-Length") or 0)
        if n > MAX_BODY:
            raise ValueError("Body zu gross")
        raw = self.rfile.read(n) if n else b""
        return json.loads(raw.decode("utf-8")) if raw.strip() else {}

    def _pin_ok(self):
        """-> None wenn ok, sonst Fehlermeldung (403)."""
        if not pin_set():
            return "pin_not_set"
        got = (self.headers.get("X-Control-Pin") or "").strip()
        if not hmac.compare_digest(got.encode(), str(CFG["control_pin"]).encode()):
            return "pin_wrong"
        return None

    # ---------------- GET (lesend, kein PIN) ----------------
    def do_GET(self):
        path = self.path.split("?", 1)[0]
        try:
            if path == "/":
                return self._send(200, ctype="text/html; charset=utf-8", body=PAGE)
            if path == "/state.json":
                c = ctx()
                live = copy.deepcopy(c.live)
                cur_live = (live.get("current") or {}).get("live") or {}
                return self._send(200, {"control": dict(c.public(), pin_default=pin_default()),
                                        "sweep": live, "matrix": c.matrix_status(),
                                        "wall_power_quelle": cur_live.get("wall_src"),
                                        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S")})
            if path == "/status.json":
                try:
                    st = read_status_once()
                except Exception as e:
                    return self._send(502, dict(exc_body(e), miner_url=sf.MINER_BASE))
                return self._send(200, {"miner_url": sf.MINER_BASE, "wall_power_w": st["_effective_wall_power_w"],
                                        "wall_power_quelle": st["_wall_power_source"], "status": st})
            if path == "/lang/list":
                return self._send(200, {"default": CFG.get("language") or "en", "languages": lang_list()})
            m = re.match(r"^/lang/([^/]+)$", path)
            if m:
                tbl = lang_table(m.group(1))
                return self._send(200, tbl) if tbl is not None else self._send(404, err("lang_not_found", "language not found"))
            if path == "/auth-status":
                return self._send(200, {"pin_set": pin_set()})   # nur ob gesetzt - niemals den PIN selbst
            if path == "/subnets":
                return self._send(200, {"subnets": local_subnets()})
            if path == "/runs":
                return self._send(200, {"runs": [run_summary(f) for f in run_files()]})
            m = re.match(r"^/runs/([^/]+)$", path)
            if m:
                d = load_run(m.group(1))
                return self._send(200, d) if d is not None else self._send(404, err("run_not_found", "Run not found"))
            m = re.match(r"^/export/([^/]+)\.(csv|json|xlsx|pdf)$", path)
            if m:
                rid, ext = m.group(1), m.group(2)
                d = load_run(rid)
                if d is None:
                    return self._send(404, err("run_not_found", "Run not found"))
                if ext == "json":
                    return self._send(200, body=json.dumps(d, indent=2, ensure_ascii=False, default=str),
                                      filename=f"{rid}.json")
                if ext == "csv":
                    return self._send(200, ctype="text/csv; charset=utf-8", body=export_csv(d), filename=f"{rid}.csv")
                try:
                    if ext == "xlsx":
                        return self._send(200, ctype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                          body=export_xlsx(d, rid), filename=f"{rid}.xlsx")
                    return self._send(200, ctype="application/pdf", body=export_pdf(d, rid), filename=f"{rid}.pdf")
                except ImportError as e:
                    return self._send(501, err("export_missing", f"Export package missing ({e}) - start control.py with the venv Python"))
            return self._send(404, err("not_found", "not found"))
        except BrokenPipeError:
            pass
        except Exception as e:
            traceback.print_exc()
            return self._send(500, {"error": f"{type(e).__name__}: {e}"})

    # ---------------- POST ----------------
    def do_POST(self):
        path = self.path.split("?", 1)[0]
        try:
            body = self._body()
        except Exception as e:
            return self._send(400, err("bad_body", f"invalid body: {e}"))
        try:
            # --- Not-Aus: bewusst OHNE PIN ---
            if path == "/emergency-stop":
                return self._emergency_stop()
            # --- PIN setzen/aendern: eigene Pruefung in _set_pin. AUSNAHME: die Ersteinrichtung
            #     (noch kein PIN gesetzt) ist neben dem Not-Aus die EINZIGE schreibende Aktion ohne PIN. ---
            if path == "/set-pin":
                return self._set_pin(body)
            # --- Miner-Suche: rein lesend, kein PIN ---
            if path == "/search":
                return self._search(body)
            # --- Ziel-Hashrate-Auswertung: liest nur runs/<id>.json, kein PIN ---
            if path == "/best-for":
                return self._best_for(body)
            # --- Plan-Generierung: rein lesend (nur GET am Miner), kein PIN ---
            if path == "/plan":
                return self._plan(body)
            # --- alle anderen schreibenden Aktionen: PIN ---
            if path in ("/connect", "/sweep/start", "/sweep/stop", "/point/set", "/mining/resume",
                        "/identify"):
                err = self._pin_ok()
                if err:
                    txt = {"pin_not_set": "Control PIN not set yet - set it in the UI under 'Set control PIN'",
                           "pin_wrong": "PIN missing or wrong (header X-Control-Pin)"}[err]
                    clog(f"403 {path}: {txt}")
                    return self._send(403, globals()["err"](err, txt))
                return getattr(self, "_post_" + path.strip("/").replace("/", "_").replace("-", "_"))(body)
            return self._send(404, err("not_found", "not found"))
        except BrokenPipeError:
            pass
        except Exception as e:
            traceback.print_exc()
            return self._send(500, {"error": f"{type(e).__name__}: {e}"})

    def _best_for(self, body):
        rid = str(body.get("run_id") or "")
        try:
            target = float(str(body.get("target_ths")).replace(",", "."))
        except (TypeError, ValueError):
            return self._send(400, err("target_required", "target_ths (number) required"))
        if not (0 < target < 10000):
            return self._send(400, err("target_required", "target_ths must be > 0"))
        d = load_run(rid)
        if d is None:
            return self._send(404, err("run_not_found", "Run not found"))
        return self._send(200, best_for(d, rid, target))

    def _set_pin(self, body):
        new_pin = body.get("new_pin")
        new_pin = new_pin.strip() if isinstance(new_pin, str) else new_pin
        with PIN_LOCK:
            first_setup = not pin_set()
            if not first_setup:
                cur = body.get("current_pin")
                cur = cur.strip() if isinstance(cur, str) else cur
                if not isinstance(cur, str) or not hmac.compare_digest(cur.encode(), str(CFG["control_pin"]).encode()):
                    time.sleep(1.0)   # Brute-Force bremsen
                    clog("403 /set-pin: current PIN missing or wrong")
                    return self._send(403, err("current_pin_wrong", "PIN already set - correct current_pin required"))
            perr = validate_new_pin(new_pin)
            if perr:
                clog(f"400 /set-pin: {perr}")
                return self._send(400, err("pin_invalid", perr))
            try:
                write_pin_to_config(new_pin)
            except Exception as e:
                clog(f"!!! /set-pin: could not write config.json ({type(e).__name__})")
                return self._send(500, err("config_write", "could not write config.json"))
        clog("Control PIN " + ("set (initial setup)" if first_setup else "changed") + " - effective immediately")
        return self._send(200, {"ok": True, "pin_set": True} if first_setup else {"ok": True})

    def _emergency_stop(self):
        c = ctx()
        with c.lock:
            c.emergency = True
            c.emergency_at = time.strftime("%Y-%m-%dT%H:%M:%S")
        c.engine_state["emergency"] = True
        c.stop_event.set()  # laufenden Sweep stoppen -> Thread setzt Tuning im finally zurueck, Mining bleibt AUS
        try:
            code, text = base.post_json(sf.MINING, {"enabled": False})
            clog(f"!!! EMERGENCY STOP: POST /mining {{\"enabled\": false}} -> HTTP {code} {text.strip()}")
            return self._send(200, {"ok": True, "http": code, "response": text.strip(), "sweep_stopped": sweep_running()})
        except Exception as e:
            clog(f"!!! EMERGENCY STOP FAILED: {e}")
            return self._send(502, {"ok": False, "error": str(e)})

    def _search(self, body):
        if not SEARCH_LOCK.acquire(blocking=False):
            return self._send(409, err("search_running", "A search is already running"))
        try:
            return self._send(200, search_miners(body.get("subnet") or "auto"))
        except ValueError as e:
            return self._send(400, exc_body(e))
        finally:
            SEARCH_LOCK.release()

    def _post_identify(self, body):
        target = body.get("miner_url_or_ip") or sf.MINER_BASE
        try:
            return self._send(200, identify_miner(target))
        except ValueError as e:
            return self._send(400, {"error": str(e)})
        except Exception as e:
            return self._send(502, {"error": f"{type(e).__name__}: {e}"})

    def _post_connect(self, body):
        url = (body.get("miner_url") or "").strip()
        if url:
            if sweep_running():
                return self._send(409, err("sweep_running_url", "Sweep running - miner URL cannot be changed"))
            try:
                url = normalize_miner_url(url)
            except ValueError as e:
                return self._send(400, err("invalid_input", f"invalid input: {e}"))
            CFG["miner_url"] = url
            apply_config()
        try:
            st, p, d = connect_and_derive()  # nur GET
        except Exception as e:
            return self._send(502, dict(exc_body(e), miner_url=sf.MINER_BASE))
        clog(f"Connected: {sf.MINER_BASE} ({p['profile']})")
        return self._send(200, {"miner_url": sf.MINER_BASE, "profile": p,
                                "wall_power": dict(zip(("w", "quelle"), sf.effective_wall_power(st))),
                                "preview": preview(p, d)})

    def _plan(self, body):
        """Testplan erzeugen: nur GET am Miner, aendert weder Miner noch globale Einstellungen."""
        over = {k: body[k] for k in ("freq_low_pct", "floor_pct", "freq_step_factor", "mv_step_factor")
                if body.get(k) not in (None, "")}
        try:
            st = read_status_once()   # nur GET
            p = profile_from(st)
            pl = sf.plan_matrix(p, body.get("sweep_mode") or "full", body.get("resolution") or "fein",
                                body.get("target_ths"), body.get("allow_above_stock"), body.get("freq_high_pct"),
                                body.get("mv_high_pct"), None, over)
        except ValueError as e:
            return self._send(400, exc_body(e))
        except Exception as e:
            return self._send(502, exc_body(e, fallback_prefix=True))
        pl["estimate"] = sf.plan_estimate(pl, p.get("local_hashrate_difficulty"))
        pl["full_measure"] = bool(body.get("full_measure", False))
        pl["profile"] = {k: p.get(k) for k in ("profile", "profile_id", "stock_frequency_mhz", "stock_core_mv",
                                                 "frequency_min", "frequency_max", "core_min", "core_max")}
        pl["created_at"] = time.strftime("%Y-%m-%dT%H:%M:%S")
        c = ctx()
        with c.lock:
            c.plan = {k: pl[k] for k in ("sweep_mode", "resolution", "band", "requires_gate", "created_at")}
        if not c.sweep_running():          # laufende Matrix nicht durch eine Vorschau ersetzen
            c.set_matrix(pl, source="plan")
        return self._send(200, pl)

    def _post_sweep_start(self, body):
        c = ctx()
        with c.lock:
            if c.sweep_running():
                return self._send(409, err("sweep_running", "A sweep is already running"))
            if c.emergency:
                return self._send(409, err("emergency_active", "Emergency stop active - use 'Resume mining' first"))
            try:
                req = self._parse_start(body)
            except ValueError as e:
                return self._send(400, {"error": str(e)})
            t = threading.Thread(target=sweep_worker, args=(req, c.id), name=f"sweep-{c.id}", daemon=False)
            c.thread = t
            c.run_info = {k: req[k] for k in ("sweep_mode", "resolution", "full_measure", "selected", "run_tag")}
            t.start()
        clog(f"Sweep start requested: {req['sweep_mode']} / {req['resolution']} / selection {req['selected'] or 'all'}")
        return self._send(202, {"ok": True, **c.run_info, "run_overrides": req["run_overrides"]})

    def _parse_start(self, body):
        mode = body.get("sweep_mode") or "efficiency"
        if mode not in sf.SWEEP_MODES:
            raise ValueError(f"sweep_mode must be one of {sf.SWEEP_MODES}")
        res = body.get("resolution") or "fein"
        if res not in sf.RESOLUTIONS:
            raise ValueError(f"resolution must be one of {sf.RESOLUTIONS}")
        target = body.get("target_ths")
        if mode == "target_hashrate" and not (isinstance(target, (int, float)) and target > 0):
            raise ValueError("target_hashrate requires target_ths > 0")
        sel = body.get("selected_points") or None
        if sel:
            sel = sorted({int(x["frequency_mhz"] if isinstance(x, dict) else x) for x in sel})
        ovr = {}
        for k, (lo, hi) in RUN_OVERRIDE_LIMITS.items():
            if body.get(k) is not None:
                v = int(body[k])
                if not lo <= v <= hi:
                    raise ValueError(f"{k} must be between {lo} and {hi}")
                ovr[k] = v
        if ovr.get("window_max_s", 10 ** 9) < ovr.get("window_min_s", 0):
            raise ValueError("window_max_s must be >= window_min_s")
        tag = body.get("run_tag")
        if tag is not None and not RUN_ID_RE.match(str(tag)):
            raise ValueError("run_tag: only letters, digits, _ and -")
        return {"sweep_mode": mode, "resolution": res, "full_measure": bool(body.get("full_measure", False)),
                "target_ths": target, "allow_above_stock": bool(body.get("allow_above_stock", False)),
                "freq_high_pct": body.get("freq_high_pct"), "confirm_gate": bool(body.get("confirm_gate", False)),
                "selected": sel, "run_overrides": ovr, "run_tag": tag, "note": str(body.get("note") or "")[:300],
                "plan_overrides": {k: body[k] for k in ("freq_low_pct", "floor_pct", "freq_step_factor", "mv_step_factor")
                                   if body.get(k) not in (None, "")}}

    def _post_sweep_stop(self, body):
        if not sweep_running():
            return self._send(409, err("no_sweep", "No sweep running"))
        ctx().stop_event.set()
        clog("Sweep stop requested -> reset to start point follows")
        return self._send(202, {"ok": True})

    def _post_point_set(self, body):
        if sweep_running():
            return self._send(409, err("sweep_running_point", "Sweep running - single point not possible"))
        try:
            freq, mv = int(body["frequency_mhz"]), int(body["core_mv"])
        except Exception:
            return self._send(400, err("point_required", "frequency_mhz and core_mv (integers) required"))
        save = bool(body.get("save", False))
        if save and body.get("confirm_save") is not True:
            return self._send(400, err("save_needs_confirm", "save=true additionally requires \"confirm_save\": true"))
        try:
            res = set_single_point(freq, mv, save)
        except sf.WallViolation as e:
            return self._send(400, exc_body(e))
        except Exception as e:
            return self._send(502, exc_body(e))
        return self._send(200, res)

    def _post_mining_resume(self, body):
        code, text = base.post_json(sf.MINING, {"enabled": True})
        c = ctx()
        with c.lock:
            c.emergency = False
        c.engine_state["emergency"] = False
        clog(f"Mining resumed: POST /mining {{\"enabled\": true}} -> HTTP {code} {text.strip()}")
        return self._send(200, {"ok": True, "http": code, "response": text.strip()})


def port_free(host, port):
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        try:
            s.bind((host, port))
            return True
        except OSError:
            return False


def lan_ip():
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("8.8.8.8", 80))
            return s.getsockname()[0]
    except OSError:
        return None


# ----------------------------------------------------------------------------
# HTML (selbstenthalten)
# ----------------------------------------------------------------------------
PAGE = r"""<!doctype html>
<html lang="en" data-theme="dark" data-accent="blue">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Miner Control Center</title>
<style>
/* ============================================================
   THEME-SYSTEM: alle Farben als Variablen
   ============================================================ */
:root, [data-theme="dark"] {
  --bg: #0a0e14;
  --bg-grad: radial-gradient(1200px 600px at 10% -10%, rgba(80,120,200,.10), transparent 60%),
             radial-gradient(900px 500px at 110% 0%, rgba(120,80,200,.08), transparent 60%);
  --card: #111823;
  --card-2: #0d131c;
  --card-border: #1f2a38;
  --card-shadow: 0 1px 0 rgba(255,255,255,.03) inset, 0 8px 24px rgba(0,0,0,.35);
  --text: #e6edf5;
  --muted: #8a99ab;
  --faint: #5d6b7c;
  --line: #213042;
  --input-bg: #0b1119;
  --input-border: #26364a;
  --chip-bg: #16202d;
  /* Status (dark) */
  --st-wait-bg: rgba(138,153,171,.10);  --st-wait-fg: #8a99ab;
  --st-act-bg: rgba(245,185,66,.16);    --st-act-fg: #f5c35a;   --st-act-solid: #f5b942;
  --st-ok-bg: rgba(52,211,130,.13);     --st-ok-fg: #4ade94;    --st-ok-solid: #22c57a;
  --st-bad-bg: rgba(248,90,90,.10);     --st-bad-fg: #f38b8b;   --st-bad-solid: #e5484d;
  --danger: #e03131; --danger-hover: #f03e3e; --danger-glow: rgba(224,49,49,.35);
}
[data-theme="light"] {
  --bg: #eef2f6;
  --bg-grad: radial-gradient(1200px 600px at 10% -10%, rgba(80,120,200,.10), transparent 60%);
  --card: #ffffff;
  --card-2: #f6f8fb;
  --card-border: #dbe3ec;
  --card-shadow: 0 1px 2px rgba(16,24,40,.05), 0 8px 24px rgba(16,24,40,.06);
  --text: #16202b;
  --muted: #5a6878;
  --faint: #8a97a6;
  --line: #e3e9f0;
  --input-bg: #ffffff;
  --input-border: #cfd8e3;
  --chip-bg: #f0f4f8;
  /* Status (light) - dunklere Vordergrundfarben fuer Lesbarkeit */
  --st-wait-bg: #eef1f5;               --st-wait-fg: #5f6b79;
  --st-act-bg: #fff4d6;                --st-act-fg: #8a5a00;   --st-act-solid: #e5a50a;
  --st-ok-bg: #e3f7ec;                 --st-ok-fg: #13733f;    --st-ok-solid: #1a9b56;
  --st-bad-bg: #fdecec;                --st-bad-fg: #b42323;   --st-bad-solid: #d63a3a;
  --danger: #d62828; --danger-hover: #e03131; --danger-glow: rgba(214,40,40,.25);
}
/* Akzentfarben */
[data-accent="blue"]   { --accent: #3b8bff; --accent-2: #1f6fe6; --accent-soft: rgba(59,139,255,.15); }
[data-accent="green"]  { --accent: #20b26b; --accent-2: #168f55; --accent-soft: rgba(32,178,107,.15); }
[data-accent="violet"] { --accent: #8b5cf6; --accent-2: #7043e0; --accent-soft: rgba(139,92,246,.16); }
[data-accent="orange"] { --accent: #f07c24; --accent-2: #d8661a; --accent-soft: rgba(240,124,36,.16); }

/* ============================================================
   GRUNDLAYOUT
   ============================================================ */
*, *::before, *::after { box-sizing: border-box; }
html, body { overflow-x: hidden; max-width: 100%; }
body {
  margin: 0; color: var(--text); background: var(--bg); background-image: var(--bg-grad); background-attachment: fixed;
  font: 14px/1.45 system-ui, -apple-system, "Segoe UI", Roboto, sans-serif; -webkit-font-smoothing: antialiased;
}
.page { max-width: 1500px; margin: 0 auto; padding: 22px clamp(14px, 3vw, 40px) 60px; }
.foot { margin-top: 28px; text-align: center; font-size: 12px; color: var(--muted); }
.foot a { color: inherit; text-decoration: none; border-bottom: 1px dotted currentColor; }
.foot a:hover { color: var(--accent); }
.foot-risk { margin-top: 6px; font-size: 11.5px; color: var(--st-act-fg); }
.stack { display: grid; gap: 18px; }
.duo { display: grid; gap: 18px; grid-template-columns: minmax(0, 5fr) minmax(0, 7fr); align-items: start; }
@media (max-width: 1320px) { .duo { grid-template-columns: minmax(0, 1fr); } }
.muted { color: var(--muted); } .faint { color: var(--faint); }
.mono { font-variant-numeric: tabular-nums; }
code { color: var(--accent); font-size: 12.5px; }

/* Karten */
.card {
  background: var(--card); border: 1px solid var(--card-border); border-radius: 14px;
  box-shadow: var(--card-shadow); min-width: 0;
}
.card-h { display: flex; align-items: center; justify-content: space-between; gap: 12px; flex-wrap: wrap;
  padding: 14px 18px 0; }
.card-h h2 { margin: 0; font-size: 13px; font-weight: 700; letter-spacing: .06em; text-transform: uppercase; color: var(--muted); }
.card-h .sub { font-size: 12.5px; color: var(--muted); }
.card-b { padding: 14px 18px 18px; min-width: 0; }
.card-b.scroll-x { overflow-x: auto; }
.subcard { background: var(--card-2); border: 1px solid var(--card-border); border-radius: 11px; padding: 12px 14px; min-width: 0; }

/* Kopfzeile */
.topbar { display: flex; align-items: center; justify-content: space-between; gap: 16px; flex-wrap: wrap; margin-bottom: 18px; }
.brand { display: flex; align-items: center; gap: 12px; min-width: 0; }
.logo { width: 38px; height: 38px; border-radius: 10px; background: linear-gradient(135deg, var(--accent), var(--accent-2));
  display: grid; place-items: center; color: #fff; font-weight: 800; box-shadow: 0 6px 18px var(--accent-soft); flex: none; }
.brand h1 { margin: 0; font-size: 20px; letter-spacing: -.01em; }
.brand .sub { font-size: 12.5px; color: var(--muted); }
.actions { display: flex; align-items: center; gap: 10px; flex-wrap: wrap; }

/* Theme-Umschalter */
.theme { display: flex; align-items: center; gap: 8px; padding: 5px; border-radius: 11px; background: var(--card); border: 1px solid var(--card-border); }
.seg { display: flex; background: var(--card-2); border-radius: 8px; padding: 2px; }
.seg button { border: 0; background: transparent; color: var(--muted); padding: 5px 10px; border-radius: 6px; font: inherit; font-size: 12.5px; cursor: pointer; }
.seg button[aria-pressed="true"] { background: var(--accent-soft); color: var(--text); font-weight: 600; }
.swatch { width: 18px; height: 18px; border-radius: 50%; border: 2px solid transparent; cursor: pointer; padding: 0; }
.swatch[aria-pressed="true"] { border-color: var(--text); box-shadow: 0 0 0 2px var(--card); }

/* Buttons & Eingaben */
.btn { border: 1px solid transparent; background: var(--accent); color: #fff; border-radius: 9px; padding: 8px 14px;
  font: inherit; font-weight: 600; cursor: pointer; white-space: nowrap; transition: filter .15s, background .15s; }
.btn:hover { filter: brightness(1.08); }
.btn.ghost { background: var(--chip-bg); color: var(--text); border-color: var(--input-border); }
.btn.small { padding: 5px 10px; font-size: 12.5px; }
.btn:disabled { opacity: .45; cursor: not-allowed; }
.estop { background: var(--danger); color: #fff; border: 0; border-radius: 12px; padding: 14px 22px; font: inherit; font-size: 17px;
  font-weight: 800; letter-spacing: .04em; cursor: pointer; box-shadow: 0 0 0 3px var(--danger-glow), 0 8px 22px var(--danger-glow); }
.estop:hover { background: var(--danger-hover); }
.estop:active { transform: translateY(1px); }
input, select { background: var(--input-bg); color: var(--text); border: 1px solid var(--input-border); border-radius: 9px;
  padding: 8px 10px; font: inherit; min-width: 0; }
input:focus, select:focus { outline: 2px solid var(--accent-soft); border-color: var(--accent); }
label.field { display: flex; flex-direction: column; gap: 4px; font-size: 12px; color: var(--muted); }
.check { display: inline-flex; align-items: center; gap: 7px; font-size: 13.5px; }
.row { display: flex; flex-wrap: wrap; gap: 10px; align-items: flex-end; }
.row > * { min-width: 0; }
.chips { display: flex; flex-wrap: wrap; gap: 8px; }
.chip { background: var(--chip-bg); border: 1px solid var(--card-border); border-radius: 999px; padding: 4px 11px; font-size: 12.5px; }
.chip b { font-weight: 650; }
.dot { display: inline-block; width: 8px; height: 8px; border-radius: 50%; background: var(--st-ok-solid); margin-right: 6px; vertical-align: 1px; box-shadow: 0 0 0 3px var(--st-ok-bg); }

/* Kacheln */
.tiles { display: grid; gap: 10px; grid-template-columns: repeat(auto-fill, minmax(128px, 1fr)); }
.tile { background: var(--card-2); border: 1px solid var(--card-border); border-radius: 11px; padding: 10px 12px; min-width: 0; }
.tile .k { font-size: 11.5px; color: var(--muted); display: flex; align-items: center; flex-wrap: wrap; gap: 4px 6px; }
.tile .v { font-size: 22px; font-weight: 700; letter-spacing: -.01em; margin-top: 3px; white-space: nowrap; font-variant-numeric: tabular-nums; }
.tile .u { font-size: 12px; color: var(--muted); margin-left: 3px; font-weight: 500; }
.tile.hl { border-color: var(--accent); box-shadow: 0 0 0 1px var(--accent) inset; }
.tile.hl .v { color: var(--accent); }
.tag { font-size: 10.5px; padding: 1px 6px; border-radius: 5px; font-weight: 600; }
.tag.est { background: var(--st-act-bg); color: var(--st-act-fg); }
.tag.meas { background: var(--st-ok-bg); color: var(--st-ok-fg); }
.phase { margin-top: 12px; display: grid; gap: 6px; }
.phase .top { display: flex; justify-content: space-between; gap: 10px; flex-wrap: wrap; font-size: 13px; }
.bar { height: 8px; background: var(--line); border-radius: 99px; overflow: hidden; }
.bar > div { height: 100%; background: var(--accent); border-radius: 99px; }
.bar.big { height: 12px; }
.bar.ok > div { background: var(--st-ok-solid); }

/* Tabellen */
table { width: 100%; border-collapse: collapse; font-variant-numeric: tabular-nums; }
th, td { padding: 8px 10px; text-align: right; white-space: nowrap; border-bottom: 1px solid var(--line); }
th { font-size: 11.5px; font-weight: 700; letter-spacing: .04em; text-transform: uppercase; color: var(--muted); background: var(--card-2); }
th:first-child, td:first-child { text-align: left; }
td.wrap, th.wrap { white-space: normal; text-align: left; min-width: 150px; }
tbody tr:last-child td { border-bottom: 0; }
.tablebox { border: 1px solid var(--card-border); border-radius: 11px; overflow: hidden; }
.tablebox .scroll { overflow-x: auto; }

/* Matrix-Status */
.pill { display: inline-block; padding: 2px 10px; border-radius: 999px; font-size: 12px; font-weight: 650; }
.pill.wait { background: var(--st-wait-bg); color: var(--st-wait-fg); }
.pill.act { background: var(--st-act-solid); color: #221700; }
.pill.ok { background: var(--st-ok-solid); color: #fff; }
.pill.bad { background: var(--st-bad-bg); color: var(--st-bad-fg); border: 1px solid var(--st-bad-fg); }
tr.s-wait td { color: var(--st-wait-fg); }
tr.s-act td { background: var(--st-act-bg); }
tr.s-ok td { background: var(--st-ok-bg); }
tr.s-bad td { background: var(--st-bad-bg); color: var(--muted); }
tr.sub td { font-size: 12.5px; padding-top: 5px; padding-bottom: 5px; }
tr.sub td:first-child { padding-left: 34px; position: relative; }
tr.sub td:first-child::before { content: ""; position: absolute; left: 18px; top: 0; bottom: 0; border-left: 2px solid var(--line); }
tr.sub.s-bad td { background: transparent; opacity: .72; }
tr.sub.s-bad td:first-child::after { content: "✕"; color: var(--st-bad-fg); margin-left: 6px; font-size: 11px; }
tr.sub.s-act td { background: var(--st-act-bg); font-weight: 650; }
tr.sub.s-act td:first-child::after { content: "●"; color: var(--st-act-solid); margin-left: 6px; animation: blink 1.2s infinite; }
tr.sub.vmin td { background: var(--st-ok-bg); color: var(--st-ok-fg); font-weight: 700; }
@keyframes blink { 50% { opacity: .25; } }
.legend { display: flex; flex-wrap: wrap; gap: 14px; font-size: 12px; color: var(--muted); }
.legend i { display: inline-block; width: 11px; height: 11px; border-radius: 3px; margin-right: 6px; vertical-align: -1px; }
.progress-head { display: flex; justify-content: space-between; align-items: baseline; gap: 10px; flex-wrap: wrap; margin-bottom: 8px; }
.progress-head b { font-size: 16px; }
.toggle { color: var(--accent); cursor: pointer; font-size: 12px; margin-left: 6px; user-select: none; }

/* Hinweise */
.note { border-radius: 10px; padding: 10px 12px; font-size: 13px; }
.note.warn { background: var(--st-act-bg); color: var(--st-act-fg); border: 1px solid color-mix(in srgb, var(--st-act-solid) 40%, transparent); }
details > summary { cursor: pointer; color: var(--muted); font-size: 13px; }
details[open] > summary { margin-bottom: 10px; }

/* Kurve */
.canvasbox { position: relative; height: 320px; }
canvas { width: 100%; height: 100%; display: block; }
.curve-legend { display: flex; flex-wrap: wrap; gap: 10px 16px; margin-top: 10px; font-size: 12px; color: var(--muted); }

/* Miner-Identitaets-Block */
.card.miner { border-left: 4px solid var(--accent); background:
  linear-gradient(90deg, var(--accent-soft), transparent 38%), var(--card); }
.miner-on { padding: 16px 20px 18px; }
.miner-head { display: flex; justify-content: space-between; align-items: flex-end; gap: 12px 20px; flex-wrap: wrap; }
.miner-kicker { display: block; font-size: 11.5px; font-weight: 700; letter-spacing: .08em; text-transform: uppercase; color: var(--accent); margin-bottom: 2px; }
.miner-name { margin: 0; font-size: 26px; font-weight: 750; letter-spacing: -.015em; line-height: 1.15; }
.miner-id { font-size: 13px; font-weight: 500; color: var(--muted); letter-spacing: 0; margin-left: 8px; }
.miner-state { display: flex; align-items: center; gap: 10px; flex-wrap: wrap; }
.conn { display: inline-flex; align-items: center; background: var(--chip-bg); border: 1px solid var(--card-border); border-radius: 999px; padding: 4px 12px; }
.conn code { font-size: 13px; }
.state-pill { font-size: 12.5px; font-weight: 700; padding: 5px 12px; border-radius: 999px; }
.state-pill.ok { background: var(--st-ok-bg); color: var(--st-ok-fg); border: 1px solid color-mix(in srgb, var(--st-ok-solid) 45%, transparent); }
.miner-facts { display: grid; grid-template-columns: repeat(auto-fit, minmax(235px, 1fr)); gap: 10px; margin-top: 14px; }
.fact { background: var(--card-2); border: 1px solid var(--card-border); border-radius: 10px; padding: 9px 13px; min-width: 0; }
.fact .fk { font-size: 11.5px; font-weight: 700; letter-spacing: .06em; text-transform: uppercase; color: var(--muted); }
.fact .fv { font-size: 17px; font-weight: 650; margin-top: 2px; font-variant-numeric: tabular-nums; overflow-wrap: anywhere; }
.fact .fv .nb { white-space: nowrap; }
.fact .fv small { font-size: 12px; font-weight: 500; color: var(--muted); }
.fact .fv .sep { color: var(--faint); margin: 0 7px; }
.fact .fv .fx { font-size: 13px; font-weight: 600; color: var(--accent); margin-left: 6px; }
.fact.now { border-color: color-mix(in srgb, var(--accent) 45%, var(--card-border)); }
.miner-off { display: none; padding: 22px 20px; color: var(--muted); font-size: 15px; }
.card.miner.disconnected { border-left-color: var(--line); background: var(--card); opacity: .75; }
.card.miner.disconnected .miner-on { display: none; }
.card.miner.disconnected .miner-off { display: block; }

/* Historie */
.runs { display: grid; gap: 14px; }
.run-h { display: flex; justify-content: space-between; align-items: center; gap: 10px; flex-wrap: wrap; margin-bottom: 10px; }
.run-h .title { font-weight: 650; }
.ss { display: grid; gap: 10px; grid-template-columns: repeat(auto-fill, minmax(220px, 1fr)); }
.ss .subcard .k { font-size: 11.5px; color: var(--muted); text-transform: uppercase; letter-spacing: .05em; }
.ss .subcard .v { font-size: 17px; font-weight: 700; margin: 2px 0; }
.ss .subcard .d { font-size: 12.5px; color: var(--muted); }
.ss .subcard.best { border-color: var(--st-ok-solid); }
.badge { font-size: 11px; padding: 2px 8px; border-radius: 6px; font-weight: 650; }
.badge.fin { background: var(--st-ok-bg); color: var(--st-ok-fg); }
.badge.short { background: var(--st-act-bg); color: var(--st-act-fg); }

/* Ergaenzungen fuer die echte Steuerzentrale */
.pinmsg { font-weight: 700; font-size: 13px; }
.pinmsg.ok { color: var(--st-ok-fg); } .pinmsg.bad { color: var(--st-bad-fg); }
.note.bad { background: var(--st-bad-bg); color: var(--st-bad-fg); border: 1px solid color-mix(in srgb, var(--st-bad-solid) 45%, transparent); }
.card.pinsetup { border-color: var(--st-act-solid); box-shadow: 0 0 0 1px var(--st-act-solid) inset, var(--card-shadow); }
.planinfo { margin: 12px 0; }
pre.log { background: var(--card-2); border: 1px solid var(--card-border); border-radius: 10px; padding: 10px 12px; max-height: 260px;
  overflow: auto; font-size: 12px; margin: 0; white-space: pre; color: var(--muted); }
details.logbox { position: relative; }
.logbtns { position: absolute; top: -3px; right: 0; display: flex; gap: 6px; }
.clip { display: inline-block; max-width: 440px; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; vertical-align: bottom; }
.pill.skip { background: var(--st-wait-bg); color: var(--st-wait-fg); opacity: .7; }
tr.s-skip td { color: var(--st-wait-fg); opacity: .6; }
tr.sub.s-pre td { color: var(--muted); }
tr.sub.s-ok td { background: var(--st-ok-bg); }
.dot.off { background: var(--st-bad-solid); box-shadow: 0 0 0 3px var(--st-bad-bg); }
.state-pill.bad { background: var(--st-bad-bg); color: var(--st-bad-fg); border: 1px solid var(--st-bad-fg); }
.badge.other { background: var(--st-wait-bg); color: var(--st-wait-fg); }
.tag:empty { display: none; }
.bestfor { margin-top: 12px; padding-top: 12px; border-top: 1px dashed var(--card-border); }
.bf-out { margin-top: 10px; display: grid; gap: 8px; }
.bf-chosen { background: var(--card); border: 1px solid var(--accent); border-radius: 11px; padding: 12px 14px; }
.bf-chosen .k { font-size: 11.5px; font-weight: 700; letter-spacing: .05em; text-transform: uppercase; color: var(--accent); }
.bf-big { font-size: 24px; font-weight: 750; margin: 2px 0 4px; }
.bf-big small { font-size: 13px; color: var(--muted); font-weight: 500; }
.bf-vals { display: flex; flex-wrap: wrap; gap: 6px 18px; font-size: 14px; }
.bf-small { font-size: 12.5px; color: var(--muted); }
.toast { position: fixed; left: 50%; bottom: 18px; transform: translateX(-50%) translateY(20px); opacity: 0; pointer-events: none;
  background: var(--card); color: var(--text); border: 1px solid var(--card-border); border-left: 4px solid var(--accent);
  border-radius: 10px; padding: 10px 16px; box-shadow: var(--card-shadow); font-weight: 600; transition: opacity .2s, transform .2s;
  max-width: calc(100vw - 32px); z-index: 50; }
.toast.show { opacity: 1; transform: translateX(-50%) translateY(0); }
.toast.bad { border-left-color: var(--st-bad-solid); color: var(--st-bad-fg); }
.langsel { padding: 4px 8px; font-size: 12.5px; border-radius: 8px; background: var(--card-2); }
</style>
</head>
<body>
<div class="page">

  <!-- HEADER -->
  <header class="topbar">
    <div class="brand">
      <div class="logo">⛏</div>
      <div>
        <h1 data-i18n="app.title">Miner Control Center</h1>
        <div class="sub"><span class="dot" id="hdot"></span><span id="hdr_profile">–</span> · <code id="miner">-</code> · <span data-i18n="hdr.sweep">Sweep</span>: <b id="sstat">-</b></div>
      </div>
    </div>
    <div class="actions">
      <div class="theme" data-i18n-title="theme.title" title="Appearance">
        <select id="langsel" class="langsel" data-i18n-title="lang.title" title="Language"></select>
        <div class="seg" role="group">
          <button type="button" data-theme-btn="dark" data-i18n="theme.dark">Dark</button>
          <button type="button" data-theme-btn="light" data-i18n="theme.light">Light</button>
        </div>
        <button type="button" class="swatch" data-accent-btn="blue"   style="background:#3b8bff" data-i18n-title="accent.blue" title="Blue"></button>
        <button type="button" class="swatch" data-accent-btn="green"  style="background:#20b26b" data-i18n-title="accent.green" title="Green"></button>
        <button type="button" class="swatch" data-accent-btn="violet" style="background:#8b5cf6" data-i18n-title="accent.violet" title="Violet"></button>
        <button type="button" class="swatch" data-accent-btn="orange" style="background:#f07c24" data-i18n-title="accent.orange" title="Orange"></button>
      </div>
      <button type="button" id="resume" class="btn ghost needpin" style="display:none" data-i18n="btn.resume">Resume mining</button>
      <button type="button" id="estop" class="estop" data-i18n="btn.estop" data-i18n-title="hint.estop">ASICs OFF NOW</button>
    </div>
  </header>
  <div id="ebanner"></div>

  <div class="stack">

    <!-- PIN SETUP -->
    <section class="card pinsetup" id="pinsetup" style="display:none">
      <div class="card-h"><h2 data-i18n="card.pinsetup.title">Set control PIN</h2><span class="sub" data-i18n="card.pinsetup.sub">write actions are locked until then – the emergency stop always works</span></div>
      <div class="card-b">
        <div class="note warn" style="margin-bottom:12px" data-i18n="hint.pin_setup">For maximum security set the PIN right after installation. At least 4 characters.</div>
        <div class="row">
          <label class="field"><span data-i18n="pin.new">New PIN</span><input id="ps_new" name="ctrl-setup-new" class="pinf" type="password" autocomplete="new-password" spellcheck="false" style="width:180px"></label>
          <label class="field"><span data-i18n="pin.repeat">Repeat new PIN</span><input id="ps_rep" name="ctrl-setup-repeat" class="pinf" type="password" autocomplete="new-password" spellcheck="false" style="width:180px"></label>
          <label class="check" style="align-self:center"><input type="checkbox" class="showpin" data-scope="ps"> &#128065; <span data-i18n="pin.show">Show PIN</span></label>
          <button type="button" id="ps_btn" class="btn" data-i18n="btn.pin_set">Set PIN</button><span id="ps_msg" class="pinmsg"></span>
        </div>
      </div>
    </section>

    <!-- CONNECTION -->
    <section class="card">
      <div class="card-h"><h2 data-i18n="card.connection.title">Connection</h2><span class="sub" id="conn_sub">–</span></div>
      <div class="card-b">
        <div class="row">
          <label class="field"><span data-i18n="conn.ip">Miner IP</span><input id="url" data-i18n-ph="conn.ip_ph" placeholder="e.g. 192.168.1.100" style="width:190px"></label>
          <button type="button" id="connect" class="btn needpin" style="display:none" data-i18n="btn.connect">Connect</button>
          <button type="button" id="search" class="btn ghost" data-i18n="btn.search">Search miner</button>
          <label class="field"><span data-i18n="conn.pin">PIN</span><input id="pin" name="ctrl-active-pin" type="password" autocomplete="new-password" spellcheck="false" style="width:130px"></label>
          <span id="pinerr" class="pinmsg" style="align-self:center"></span>
          <span class="muted" id="pinhint" style="align-self:center;font-size:12.5px"></span>
        </div>
        <details id="adv" style="margin-top:12px"><summary data-i18n="conn.advanced">Advanced: choose subnet</summary>
          <div class="row"><select id="subnet" style="min-width:280px"></select><button type="button" id="search2" class="btn ghost" data-i18n="btn.scan">Scan</button>
          <span class="muted" id="subinfo" style="align-self:center"></span></div>
          <div class="row" style="margin-top:8px"><label class="field"><span data-i18n="subnet.custom_label">Scan custom subnet</span>
            <input id="subnet_custom" data-i18n-ph="subnet.custom_ph" placeholder="e.g. 192.168.0.0/24" spellcheck="false" style="width:200px"></label>
            <button type="button" id="search3" class="btn ghost" data-i18n="btn.scan" style="align-self:flex-end">Scan</button></div>
          <div class="muted" style="font-size:12.5px;margin-top:4px" data-i18n="subnet.custom_hint">Also for networks this computer is not attached to, if reachable via your router. Private ranges only, one /24 at a time.</div></details>
        <div id="searchinfo" class="muted" style="margin-top:8px;font-size:13px"></div>
        <div class="tablebox" id="hitsbox" style="margin-top:8px;display:none"><div class="scroll"><table id="hits"></table></div></div>
        <details id="pinchange" style="display:none;margin-top:12px"><summary data-i18n="pin.change_title">Change PIN</summary>
          <div class="row">
            <label class="field"><span data-i18n="pin.current">Current PIN</span><input id="pc_cur" name="ctrl-change-current" class="pinf" type="password" autocomplete="new-password" spellcheck="false" style="width:160px"></label>
            <label class="field"><span data-i18n="pin.new">New PIN</span><input id="pc_new" name="ctrl-change-new" class="pinf" type="password" autocomplete="new-password" spellcheck="false" style="width:160px"></label>
            <label class="field"><span data-i18n="pin.repeat">Repeat new PIN</span><input id="pc_rep" name="ctrl-change-repeat" class="pinf" type="password" autocomplete="new-password" spellcheck="false" style="width:160px"></label>
            <label class="check" style="align-self:center"><input type="checkbox" class="showpin" data-scope="pc"> &#128065; <span data-i18n="pin.show">Show PIN</span></label>
            <button type="button" id="pc_btn" class="btn ghost" data-i18n="btn.pin_change">Change PIN</button><span id="pc_msg" class="pinmsg"></span>
          </div></details>
      </div>
    </section>

    <!-- MINER -->
    <section class="card miner disconnected" id="minerCard">
      <div class="miner-on">
        <div class="miner-head">
          <div class="miner-title">
            <span class="miner-kicker" data-i18n="card.miner.kicker">Miner</span>
            <h2 class="miner-name"><span id="m_name">–</span> <span class="miner-id" id="m_id"></span></h2>
          </div>
          <div class="miner-state">
            <span class="conn"><span class="dot" id="m_dot"></span><code id="m_ip">–</code></span>
            <span class="state-pill ok" id="m_state">–</span>
          </div>
        </div>
        <div class="miner-facts">
          <div class="fact"><div class="fk" data-i18n="miner.stock">Stock</div><div class="fv" id="f_stock">–</div></div>
          <div class="fact"><div class="fk" data-i18n="miner.limits">Limits</div><div class="fv" id="f_limits">–</div></div>
          <div class="fact now"><div class="fk" data-i18n="miner.current">Current</div><div class="fv"><span class="nb" id="f_now">–</span> <span class="fx nb" id="f_ths"></span></div></div>
          <div class="fact"><div class="fk" data-i18n="miner.input">Input</div><div class="fv" id="f_vin">–</div></div>
        </div>
      </div>
      <div class="miner-off" id="m_off"></div>
    </section>

    <!-- LIVE + TEST MATRIX -->
    <div class="duo">
      <section class="card">
        <div class="card-h"><h2 data-i18n="card.live.title">Live</h2><span class="sub" id="live_sub">–</span></div>
        <div class="card-b">
          <div class="tiles">
            <div class="tile"><div class="k" data-i18n="tile.freq">Frequency</div><div class="v"><span id="t_f">–</span><span class="u">MHz</span></div></div>
            <div class="tile"><div class="k" data-i18n="tile.mv">core_mv</div><div class="v"><span id="t_mv">–</span><span class="u">mV</span></div></div>
            <div class="tile"><div class="k" data-i18n="tile.hashrate">Hashrate local</div><div class="v"><span id="t_h">–</span><span class="u">TH/s</span></div></div>
            <div class="tile"><div class="k"><span data-i18n="tile.wall">Wall power</span> <span class="tag" id="t_wsrc"></span></div><div class="v"><span id="t_w">–</span><span class="u">W</span></div></div>
            <div class="tile hl"><div class="k" id="t_jk">J/TH</div><div class="v"><span id="t_j">–</span></div></div>
            <div class="tile"><div class="k" data-i18n="tile.error">Error rate</div><div class="v"><span id="t_e">–</span><span class="u">%</span></div></div>
            <div class="tile"><div class="k" data-i18n="tile.temp">ASIC / VR</div><div class="v"><span id="t_t">–</span><span class="u">°C</span></div></div>
            <div class="tile"><div class="k" data-i18n="tile.input">Input</div><div class="v"><span id="t_v">–</span><span class="u">V</span></div></div>
          </div>
          <div class="phase">
            <div class="top"><span><b id="t_ph">–</b></span><span class="muted mono" id="t_pt"></span></div>
            <div class="bar"><div id="bar" style="width:0"></div></div>
          </div>
          <details class="logbox" style="margin-top:12px"><summary data-i18n="live.log">Log</summary>
            <div class="logbtns"><button type="button" id="logcopy" class="btn ghost small" data-i18n="log.copy">Copy</button>
            <button type="button" id="logclear" class="btn ghost small" data-i18n="log.clear" data-i18n-title="log.clear_hint" title="clears the display only, not the server log">Clear</button></div>
            <pre id="log" class="log"></pre></details>
        </div>
      </section>

      <section class="card">
        <div class="card-h"><h2 data-i18n="card.matrix.title">Test matrix</h2><span class="sub" id="mx_sub"></span></div>
        <div class="card-b">
          <div class="progress-head"><b id="mx_text"></b><span class="muted" id="mx_rest"></span></div>
          <div class="bar big ok"><div id="mx_bar" style="width:0"></div></div>
          <div class="legend" style="margin:10px 0 12px">
            <span><i style="background:var(--st-wait-fg);opacity:.5"></i><span data-i18n="legend.wait">waiting</span></span>
            <span><i style="background:var(--st-act-solid)"></i><span data-i18n="legend.active">active</span></span>
            <span><i style="background:var(--st-ok-solid)"></i><span data-i18n="legend.done">done / valid</span></span>
            <span><i style="background:var(--st-bad-fg);opacity:.6"></i><span data-i18n="legend.rejected">rejected / underpowered</span></span>
            <span><i style="background:var(--st-ok-bg);border:1px solid var(--st-ok-fg)"></i><span data-i18n="legend.vmin">result Vmin</span></span>
            <span id="mx_early"></span>
          </div>
          <div class="tablebox"><div class="scroll"><table id="mxtab"></table></div></div>
          <details style="margin-top:12px"><summary data-i18n="matrix.all_points">All measurement points (detail table)</summary>
            <div class="tablebox"><div class="scroll"><table id="restab"></table></div></div></details>
        </div>
      </section>
    </div>

    <!-- PLAN SWEEP -->
    <section class="card">
      <div class="card-h"><h2 data-i18n="card.plan.title">Plan sweep</h2><span class="sub" data-i18n="card.plan.sub">Plan preview · no writes until “Start sweep”</span></div>
      <div class="card-b">
        <div class="row">
          <label class="field"><span data-i18n="plan.goal">Test goal</span>
            <select id="p_mode">
              <option value="efficiency" data-i18n="mode.efficiency">Efficiency (lower range)</option>
              <option value="performance" data-i18n="mode.performance">Max performance (upper range)</option>
              <option value="full" selected data-i18n="mode.full">Complete</option>
              <option value="target_hashrate" data-i18n="mode.target_hashrate">Target hashrate</option></select></label>
          <label class="field" id="l_target" style="display:none"><span data-i18n="plan.target_ths">Target TH/s</span><input id="p_target" type="number" step="0.1" min="0.1" value="5" style="width:90px"></label>
          <label class="field"><span data-i18n="plan.resolution">Resolution</span><select id="p_res"><option value="fein" data-i18n="res.fine">fine</option><option value="grob" data-i18n="res.coarse">coarse</option></select></label>
          <label class="check" data-i18n-title="hint.full_measure"><input id="p_full" type="checkbox"> <span data-i18n="plan.full_measure">Measure every point fully</span></label>
          <label class="check"><input id="p_above" type="checkbox"> <span data-i18n="plan.above_stock">Allow above stock</span></label>
          <label class="field" id="l_fhp" style="display:none"><span data-i18n="plan.freq_plus">+ frequency</span><select id="p_fhp"><option value="0.10">10 %</option><option value="0.15">15 %</option><option value="0.20">20 %</option></select></label>
          <div style="margin-left:auto" class="row">
            <button type="button" id="plan" class="btn ghost" data-i18n="btn.plan">Create plan</button>
            <button type="button" id="start" class="btn needpin" style="display:none" disabled data-i18n="btn.start" data-i18n-title="hint.start_needs_plan">Start sweep</button>
            <button type="button" id="stop" class="btn ghost needpin" style="display:none" disabled data-i18n="btn.stop">Stop sweep</button>
          </div>
        </div>
        <div id="abovewarn" class="note warn" style="display:none;margin-top:12px" data-i18n="hint.above_stock">Caution: above stock the test leaves the manufacturer's factory range (higher load, heat, wear). The hard wall frequency_max/core_max always stays active; a confirmation is required at start.</div>
        <details style="margin-top:12px">
          <summary data-i18n="plan.fine_params">Fine parameters (optional, only for this plan)</summary>
          <div class="row">
            <label class="field"><span data-i18n="plan.band_low">Band low</span><input id="p_freq_low_pct" type="number" step="0.01" min="0" max="0.95" style="width:90px"></label>
            <label class="field"><span data-i18n="plan.floor">Floor</span><input id="p_floor_pct" type="number" step="0.01" min="0.3" max="1" style="width:90px"></label>
            <label class="field"><span data-i18n="plan.fstep">f step ×</span><input id="p_freq_step_factor" type="number" step="1" min="1" placeholder="25" style="width:80px"></label>
            <label class="field"><span data-i18n="plan.mvstep">mV step ×</span><input id="p_mv_step_factor" type="number" step="1" min="1" placeholder="5" style="width:80px"></label>
          </div>
        </details>
        <div id="gate" style="margin-top:12px"></div>
        <div id="planinfo" class="planinfo"></div>
        <div class="tablebox" id="planbox" style="display:none"><div class="scroll"><table id="plantab"></table></div></div>
      </div>
    </section>

    <!-- CURVE -->
    <section class="card">
      <div class="card-h"><h2 data-i18n="card.curve.title">J/TH curve</h2><span class="sub" data-i18n="card.curve.sub">valid points of the current run · colour = frequency</span></div>
      <div class="card-b">
        <div class="canvasbox"><canvas id="chart"></canvas></div>
        <div class="curve-legend" id="curveLegend"></div>
      </div>
    </section>

    <!-- HISTORY -->
    <section class="card">
      <div class="card-h"><h2 data-i18n="card.history.title">History</h2><span class="sub" id="hist_sub"></span></div>
      <div class="card-b"><div class="runs" id="hist"></div></div>
    </section>

  </div>
  <footer class="foot"><a href="https://github.com/NFDiJee/minertune" target="_blank" rel="noopener" data-i18n-title="footer.source">MinerTune</a>
    · <span data-i18n="footer.made_by">made by NFDiJee</span>
    · <a href="https://github.com/NFDiJee/minertune/blob/main/LICENSE" target="_blank" rel="noopener" data-i18n="footer.license">MIT License</a>
    <div class="foot-risk">⚠ <b data-i18n="risk.title">Use at your own risk</b> – <span data-i18n="risk.text">MinerTune changes voltage/clock settings – incorrect use can damage your miner. No warranty.</span></div></footer>
</div>
<div id="toast" class="toast"></div>
<script>
const $=id=>document.getElementById(id);
const root=document.documentElement;
function lsGet(k){try{return localStorage.getItem(k)||""}catch(e){return ""}}
function lsSet(k,v){try{localStorage.setItem(k,v)}catch(e){}}
const esc=s=>String(s??"").replace(/[&<>"]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
function hostOf(u){try{return new URL(u).host}catch(e){return u||""}}

/* ================= i18n ================= */
let EN={}, LANG={}, LANGCODE="en", LANGS=[];
/* t(): gewaehlte Sprache -> Fallback Englisch (Referenz) -> nie roher Schluessel */
function t(key,vars){let s=LANG[key];if(s===undefined||s==="")s=EN[key];
  if(s===undefined||s===""){s=key.split(".").pop().replace(/_/g," ")}
  if(vars)s=s.replace(/\{(\w+)\}/g,(m,k)=>vars[k]!==undefined?vars[k]:m);return s}
const DEC=()=>LANG._decimal||EN._decimal||".";
const fmt=(v,d)=>(v===null||v===undefined||v===""||Number.isNaN(Number(v)))?"–":Number(v).toFixed(d).replace(".",DEC());
const F={w:v=>fmt(v,1),ths:v=>fmt(v,3),jth:v=>fmt(v,2),volt:v=>fmt(v,2),temp:v=>fmt(v,1),pct:v=>fmt(v,2),int:v=>fmt(v,0),
  num:v=>(v===null||v===undefined)?"–":String(v).replace(".",DEC())};
function dur(s){const h=Math.floor(s/3600),m=Math.round((s%3600)/60);return h?t("unit.h_min",{h,m}):t("unit.min",{m})}
/* Engine-Meldungen {code, params} -> Schluessel code.<code>; Listen mit " · "; alte Freitexte (Altlaeufe) unveraendert */
function trMsg(x){if(x===null||x===undefined||x==="")return "";
  if(Array.isArray(x))return x.map(trMsg).filter(Boolean).join(" · ");
  if(typeof x==="object"&&x.code){const P={};
    for(const [k,v] of Object.entries(x.params||{}))P[k]=typeof v==="number"?F.num(v):(k==="sensor"?String(v).toUpperCase():v);
    return hasKey("code."+x.code)?t("code."+x.code,P):x.code}
  return String(x)}
const hasKey=k=>(LANG[k]!==undefined&&LANG[k]!=="")||(EN[k]!==undefined&&EN[k]!=="");
function trErr(e){if(!e)return "";if(e.code){if(hasKey("err."+e.code))return t("err."+e.code,e.params||{});
    if(hasKey("code."+e.code))return trMsg({code:e.code,params:e.params})}
  return String(e.message||"")}

/* Statische Texte (data-i18n / -ph / -title) */
function applyStatic(){root.lang=LANGCODE;document.title=t("app.title");
  document.querySelectorAll("[data-i18n]").forEach(e=>e.textContent=t(e.dataset.i18n));
  document.querySelectorAll("[data-i18n-ph]").forEach(e=>e.placeholder=t(e.dataset.i18nPh));
  document.querySelectorAll("[data-i18n-title]").forEach(e=>e.title=t(e.dataset.i18nTitle));
  $("p_freq_low_pct").placeholder=F.num(0.52);$("p_floor_pct").placeholder=F.num(0.75);
  $("m_off").innerHTML=t("miner.not_connected");}
async function loadLang(code){const r=await fetch("/lang/"+encodeURIComponent(code),{cache:"no-store"});if(!r.ok)throw new Error("lang");return r.json()}
async function setLang(code,persist){try{LANG=code==="en"?EN:await loadLang(code);LANGCODE=code}catch(e){LANG=EN;LANGCODE="en"}
  if(persist)lsSet("ctrl_lang",LANGCODE);$("langsel").value=LANGCODE;applyStatic();
  if(last)render(last);if(lastPlan)showModePlan(lastPlan);loadHist();if(lastIdle)showIdle(lastIdle);refreshAuth();redrawChart();}
async function initLang(){try{const r=await (await fetch("/lang/list",{cache:"no-store"})).json();LANGS=r.languages||[];
  EN=await loadLang("en");
  $("langsel").innerHTML=LANGS.map(l=>`<option value="${esc(l.code)}">${esc(l.name)}</option>`).join("");
  const pref=lsGet("ctrl_lang"),code=LANGS.some(l=>l.code===pref)?pref:(LANGS.some(l=>l.code===r.default)?r.default:"en");
  await setLang(code,false)}catch(e){applyStatic()}}
$("langsel").onchange=()=>setLang($("langsel").value,true);

/* ================= Theme / Akzent ================= */
function applyTheme(t_){root.dataset.theme=t_;document.querySelectorAll("[data-theme-btn]").forEach(b=>b.setAttribute("aria-pressed",b.dataset.themeBtn===t_));lsSet("ctrl_theme",t_);redrawChart()}
function applyAccent(a){root.dataset.accent=a;document.querySelectorAll("[data-accent-btn]").forEach(b=>b.setAttribute("aria-pressed",b.dataset.accentBtn===a));lsSet("ctrl_accent",a);redrawChart()}
document.querySelectorAll("[data-theme-btn]").forEach(b=>b.onclick=()=>applyTheme(b.dataset.themeBtn));
document.querySelectorAll("[data-accent-btn]").forEach(b=>b.onclick=()=>applyAccent(b.dataset.accentBtn));

/* ================= API / PIN ================= */
$("pin").value=lsGet("ctrl_pin").trim();
$("pin").addEventListener("input",()=>{lsSet("ctrl_pin",$("pin").value.trim());$("pinerr").textContent=""});
function activePin(){return $("pin").value.trim()}
async function api(method,path,body,needPin){
  const h={"Content-Type":"application/json"};if(needPin)h["X-Control-Pin"]=activePin();
  const r=await fetch(path,{method,headers:h,body:body?JSON.stringify(body):undefined,cache:"no-store"});
  let j={};try{j=await r.json()}catch(e){}
  if(needPin&&r.status===403){const e=$("pinerr");e.className="pinmsg bad";e.textContent=t("msg.pin_wrong")}
  else if(needPin&&r.ok)$("pinerr").textContent="";
  if(!r.ok){const err=new Error((j&&j.error)||("HTTP "+r.status));err.status=r.status;err.code=j&&j.code;err.params=j&&j.params;throw err}return j;}
let toastT=null;
function toast(msg,bad){const x=$("toast");x.textContent=msg;x.className="toast show"+(bad?" bad":"");clearTimeout(toastT);toastT=setTimeout(()=>x.className="toast",bad?7000:4000)}

/* ================= Not-Aus / Mining fortsetzen ================= */
$("estop").onclick=async()=>{try{await api("POST","/emergency-stop",{},false);toast(t("msg.estop_sent"),true)}catch(e){toast(t("msg.estop_failed",{why:trErr(e)}),true)}};
$("resume").onclick=async()=>{try{await api("POST","/mining/resume",{},true);toast(t("msg.resumed"))}catch(e){toast(trErr(e),true)}};

/* ================= Miner suchen / Subnetze / Identify ================= */
async function doSearch(subnet){$("searchinfo").textContent=t("search.running",{subnet});$("hits").innerHTML="";$("hitsbox").style.display="none";
  try{const j=await api("POST","/search",{subnet},false);
   $("searchinfo").textContent=t("search.result",{n:j.hits.length,nets:j.subnets.join(", "),scanned:j.scanned,s:F.num(j.duration_s)});
   if(j.hits.length){$("hitsbox").style.display="";
    $("hits").innerHTML=`<thead><tr><th>${t("col.profile")}</th><th>${t("col.ip")}</th><th>${t("col.op_point")}</th><th>TH/s</th><th>${t("col.mining")}</th><th>${t("col.temp")}</th><th></th></tr></thead><tbody>`+
    j.hits.map(h=>`<tr><td><b>${esc(h.profile)}</b> <span class=muted>${esc(h.profile_id)}</span></td><td><code>${esc(h.ip)}</code></td><td>${h.current_frequency_mhz} MHz / ${h.core_mv} mV</td>
    <td>${F.ths(h.local_hashrate_ths)}</td><td>${h.mining_enabled?t("miner.mining_on_short"):`<b style='color:var(--st-bad-fg)'>${t("miner.mining_off_short")}</b>`}</td><td>${F.temp(h.asic_temp_c)} / ${F.temp(h.vr_temp_c)}</td>
    <td><button type="button" class="btn ghost small" data-sel="${esc(h.ip)}">${t("btn.select")}</button> <button type="button" class="btn ghost small" data-id="${esc(h.ip)}" style="${pinIsSet?"":"display:none"}">${t("btn.identify")}</button></td></tr>`).join("")+"</tbody>";}
   document.querySelectorAll("[data-sel]").forEach(b=>b.onclick=()=>{$("url").value=b.dataset.sel;$("connect").click()});
   document.querySelectorAll("[data-id]").forEach(b=>b.onclick=async()=>{b.disabled=true;try{const r=await api("POST","/identify",{miner_url_or_ip:b.dataset.id},true);
     $("searchinfo").textContent=r.identified?t("search.identified",{ip:b.dataset.id}):t("err."+(r.code||"no_display"))}catch(e){$("searchinfo").textContent=t("btn.identify")+": "+trErr(e)}b.disabled=false});
  }catch(e){$("searchinfo").textContent=t("search.failed",{why:trErr(e)})}}
$("search").onclick=()=>doSearch("auto");
$("search2").onclick=()=>doSearch($("subnet").value);
$("search3").onclick=()=>{const v=$("subnet_custom").value.trim();if(!v){$("searchinfo").textContent=t("err.subnet_invalid");return}doSearch(v)};
$("subnet_custom").addEventListener("keydown",e=>{if(e.key==="Enter"){e.preventDefault();$("search3").click()}});
const SUBLBL={"aktiv/LAN":"subnet.lan","docker":"subnet.docker","vpn/tailscale":"subnet.vpn","sonstiges":"subnet.other"};
async function fillSubnets(){try{const j=await api("GET","/subnets");
  $("subnet").innerHTML=j.subnets.map(s=>`<option value="${esc(s.subnet)}"${s.recommended?" selected":""}>${esc(s.subnet)} – ${esc(SUBLBL[s.label]?t(SUBLBL[s.label]):s.label)} (${esc(s.interface)})${s.recommended?" – "+t("subnet.recommended"):""}</option>`).join("")+`<option value="all">${t("subnet.all")}</option>`;
  $("subinfo").textContent=t("subnet.found",{n:j.subnets.length})}catch(e){$("subinfo").textContent=trErr(e)}}
$("adv").addEventListener("toggle",()=>{if($("adv").open)fillSubnets()});

/* ================= Verbinden + MINER-Block ================= */
let minerSeen=null,minerData=null;
$("connect").onclick=async()=>{try{const j=await api("POST","/connect",{miner_url:$("url").value.trim()},true);
  fillMiner(j.profile,j.miner_url);showPlanPreview(j.preview);toast(t("msg.connected",{name:j.profile.profile||hostOf(j.miner_url)}))}catch(e){toast(trErr(e),true)}};
function fillMiner(p,url){if(!p||p.stock_frequency_mhz==null){setMinerOff(t("miner.not_connected"));return}
  minerSeen=minerSeen||new Date();minerData={p,url};$("minerCard").classList.remove("disconnected");
  $("m_name").textContent=p.profile||t("miner.unknown");$("m_id").textContent=p.profile_id||"";$("m_ip").textContent=hostOf(url||(last&&last.control.miner_url)||"");
  const on=p.mining_enabled!==false;const st=$("m_state");st.textContent=on?t("miner.mining_on"):t("miner.mining_off");st.className="state-pill "+(on?"ok":"bad");
  $("m_dot").className="dot"+(on?"":" off");
  $("f_stock").innerHTML=`${p.stock_frequency_mhz} <small>MHz</small> / ${p.stock_core_mv} <small>mV</small>`;
  $("f_limits").innerHTML=`<span class="nb">${p.frequency_min}–${p.frequency_max} <small>MHz</small></span><span class="sep">·</span><span class="nb">${p.core_min}–${p.core_max} <small>mV</small></span>`;
  $("f_now").innerHTML=`${p.current_frequency_mhz} <small>MHz</small> / ${p.core_mv} <small>mV</small>`;
  $("f_ths").textContent=p.local_hashrate_ghs!=null?F.ths(p.local_hashrate_ghs/1000)+" TH/s":"";
  $("f_vin").innerHTML=p.input_voltage_v!=null?`${F.volt(p.input_voltage_v)} <small>V</small>`:"–";
  $("hdr_profile").textContent=p.profile||"–";$("hdot").className="dot"+(on?"":" off");}
let minerOffWhy=null;
function setMinerOff(html){$("minerCard").classList.add("disconnected");$("m_off").innerHTML=html;$("hdot").className="dot off"}

/* ================= Sweep planen / starten / stoppen ================= */
$("p_mode").onchange=()=>{$("l_target").style.display=$("p_mode").value==="target_hashrate"?"":"none"};
$("p_above").onchange=()=>{const on=$("p_above").checked;$("l_fhp").style.display=on?"":"none";$("abovewarn").style.display=on?"block":"none"};
let lastPlan=null,lastPlanReq=null;
$("plan").onclick=async()=>{const b={sweep_mode:$("p_mode").value,resolution:$("p_res").value,allow_above_stock:$("p_above").checked,full_measure:$("p_full").checked};
  if(b.allow_above_stock)b.freq_high_pct=Number($("p_fhp").value);if(b.sweep_mode==="target_hashrate")b.target_ths=Number($("p_target").value);
  for(const k of["freq_low_pct","floor_pct","freq_step_factor","mv_step_factor"]){const v=$("p_"+k).value;if(v!=="")b[k]=Number(v)}
  try{showModePlan(await api("POST","/plan",b,false));lastPlanReq=b}catch(e){toast(trErr(e),true)}};
function chip(html){return `<span class="chip">${html}</span>`}
function bandText(pl){return trMsg(pl.band&&pl.band.reason)}
function showModePlan(pl){const E=pl.estimate,T=pl.target;lastPlan=pl;
  $("gate").innerHTML=pl.requires_gate?`<div class="note bad">${t("hint.gate_prefix")} ${esc(trMsg(pl.gate_text))}</div>`:"";
  const c=[chip(`<b>${t("mode."+pl.sweep_mode)}</b> · ${t(pl.resolution==="grob"?"res.coarse":"res.fine")}`),
    chip(t("plan.chip.band",{lo:pl.band.from_mhz,hi:pl.band.to_mhz,step:pl.freq_step})),
    chip(t("plan.chip.voltage",{a:pl.mv_stock+(pl.mv_top!==pl.mv_stock?"/"+pl.mv_top:""),b:pl.mv_floor,step:pl.mv_step})),
    chip(t("plan.chip.order",{o:t(isAsc(pl.order)?"order.asc":"order.desc")})),
    chip(t("plan.chip.realistic",{fast:E.fast?dur(E.fast.duration_realistic_s):"–",full:dur(E.full_measure.duration_realistic_s)})+(pl.full_measure?" · "+t("plan.chosen_full"):E.fast?" · "+t("plan.chosen_fast"):"")),
    chip(t("plan.chip.worst",{fast:E.fast?dur(E.fast.duration_worst_s):"–",full:dur(E.full_measure.duration_worst_s)})),chip(t("plan.chip.cal",{n:E.calibration_points}))];
  if(pl.anchors.length)c.push(chip(t("plan.chip.anchors",{list:pl.anchors.join(", ")})));
  if(T)c.push(chip(t("plan.chip.target",{t:F.num(T.target_ths),f:F.num(T.f_target_mhz)})+(T.preliminary?" ("+t("plan.preliminary")+")":"")));
  if(pl.early_stop_rule)c.push(chip(t("plan.chip.early_stop")));
  const notes=[esc(bandText(pl))];if(E.fast)notes.push(esc(trMsg(E.fast.note)));
  if(pl.early_stop_rule)notes.push(esc(trMsg(pl.early_stop_rule)));pl.notes.forEach(n=>notes.push("<b>"+t("word.note")+":</b> "+esc(trMsg(n))));
  $("planinfo").innerHTML=`<div class="chips">${c.join("")}</div><div class="muted" style="font-size:12.5px;margin-top:8px">${notes.join(" · ")}</div>`;
  $("planbox").style.display="";
  $("plantab").innerHTML=`<thead><tr><th><input type=checkbox id=pall checked> #</th><th>${t("col.frequency")}</th><th>${t("col.role")}</th><th>${t("col.mv_range")}</th><th>${t("col.steps")}</th><th>${t("col.exp_ths")}</th><th>${t("col.window")}</th><th class=wrap>${t("col.hint")}</th></tr></thead><tbody>`+
   pl.points.map(q=>`<tr><td><label class="check"><input type=checkbox class=pf value=${q.freq} checked> ${q.index}</label></td><td><b>${q.freq} MHz</b></td><td>${q.role==="anchor"?"<b>"+t("role.anchor")+"</b>":t("role.band")}</td>
   <td class="mono">${q.mv_start} ${isAsc(pl.order)?"↑":"↓"} ${q.mv_floor}</td><td>${q.steps}</td><td>${fmt(q.expected_ths,2)}</td><td>${q.window_s} s</td>
   <td class="wrap">${q.above_stock?"<b style='color:var(--st-bad-fg)'>"+t("plan.above_stock_flag")+"</b> ":""}${q.capped_to_wall?"<b>"+t("plan.capped_flag")+"</b>":""}</td></tr>`).join("")+"</tbody>";
  $("pall").onchange=e=>document.querySelectorAll(".pf").forEach(x=>x.checked=e.target.checked);}
let lastPreview=null;
function showPlanPreview(pv){if(!pv)return;lastPreview=pv;const d=pv.derived;
  $("planinfo").innerHTML=`<div class="chips">${chip(t("plan.preview.band",{lo:d.f_low,hi:d.f_high,step:d.freq_sweep_step}))}${chip(t("plan.preview.voltage",{a:d.mv_start,b:d.mv_floor}))}${chip(t("plan.preview.start",{f:d.safe[0],mv:d.safe[1]}))}${chip(t("plan.preview.next"))}</div>`;}
$("start").onclick=async()=>{if(!lastPlan){toast(t("msg.plan_first"),true);return}
  const sel=[...new Set([...document.querySelectorAll(".pf:checked")].map(c=>Number(c.value)))];
  if(!sel.length){toast(t("msg.no_freq"),true);return}
  const b=Object.assign({},lastPlanReq,{selected_points:sel});
  let msg=t("confirm.start",{mode:t("mode."+lastPlan.sweep_mode),res:t(lastPlan.resolution==="grob"?"res.coarse":"res.fine"),n:sel.length})+(b.full_measure?"\n"+t("confirm.start_full"):"");
  if(lastPlan.requires_gate){msg+="\n\n"+t("confirm.gate");b.confirm_gate=true}
  if(!confirm(msg))return;
  try{await api("POST","/sweep/start",b,true);toast(t("msg.sweep_started"))}catch(e){toast(trErr(e),true)}};
$("stop").onclick=async()=>{try{await api("POST","/sweep/stop",{},true);toast(t("msg.stop_requested"))}catch(e){toast(trErr(e),true)}};

/* ================= Test-Matrix ================= */
const mxOpen={};
const ROWCLS={wartet:"wait",aktiv:"act",fertig:"ok",verworfen:"bad",uebersprungen:"skip"};
const ROWKEY={wartet:"status.waiting",aktiv:"status.active",fertig:"status.done",verworfen:"status.rejected",uebersprungen:"status.skipped"};
const SUBCLS={gueltig:"ok",vorab_ok:"pre",verworfen:"bad",aktiv:"act"};
const TAGKEY={reserve:"tag.reserve",probe_down:"tag.probe_down",knee_reserve:"tag.knee_reserve",precheck:"tag.precheck"};
const isAsc=o=>o==="asc"||o==="aufsteigend";   /* "aufsteigend": Altlaeufe */
const phase=p=>p?t("phase."+p):"–";
function renderMatrix(M){
  if(!M){$("mx_text").textContent=t("matrix.no_plan");$("mx_rest").textContent="";$("mx_sub").textContent="";$("mx_bar").style.width="0";$("mx_early").textContent="";
    $("mxtab").innerHTML=`<tbody><tr><td class='muted'>${t("matrix.empty_hint")}</td></tr></tbody>`;return}
  const P=M.progress;
  $("mx_text").textContent=(M.source==="plan"?t("matrix.preview_prefix")+" ":"")+(P.calibrating?t("matrix.calibrating"):t("matrix.progress",{x:P.x,n:P.n}))+(!M.running&&P.done&&M.source==="sweep"?" · "+t("matrix.run_done"):"");
  $("mx_sub").textContent=[M.sweep_mode?t("mode."+M.sweep_mode):null,M.order?t(isAsc(M.order)?"order.asc":"order.desc"):null].filter(Boolean).join(" · ");
  $("mx_rest").textContent=P.remaining_s!=null?t("matrix.remaining",{d:dur(P.remaining_s)}):"";
  $("mx_bar").style.width=(P.n?100*P.done/P.n:0)+"%";
  $("mx_early").textContent=M.early_stop?t("matrix.early_stop",{after:trMsg(M.early_stop)}):"";
  let h=`<thead><tr><th>${t("col.frequency")}</th><th>${t("col.status")}</th><th>${t("col.search")}</th><th>Vmin</th><th>TH/s</th><th>J/TH</th><th class=wrap>${t("col.info")}</th></tr></thead><tbody>`;
  for(const r of M.rows){const cls=ROWCLS[r.status]||"wait",open=r.status==="aktiv"||mxOpen[r.freq],n=r.sub.length;let info="";
    if(r.status==="aktiv"){const a=r.sub.find(x=>x.status==="aktiv")||{};info=t("matrix.testing",{mv:a.mv,phase:phase(a.phase),s:F.int(a.elapsed_s)})}
    else if(r.status==="fertig")info=r.found!=null&&r.found!==r.vmin?t("matrix.info_found",{mv:r.found}):"";
    else if(r.status==="verworfen"){const l=r.sub[r.sub.length-1];info=l?`<span class=clip>${esc(trMsg(l.reason))}</span>`:""}
    else if(r.role==="anchor")info=t("matrix.anchor");
    h+=`<tr class="s-${cls}"><td><b>${r.freq} MHz</b>${n&&r.status!=="aktiv"?`<span class="toggle" data-f="${r.freq}">${open?"▾":"▸"} ${n}</span>`:""}</td>
     <td><span class="pill ${cls}">${t(ROWKEY[r.status]||"status.waiting")}</span></td>
     <td class="mono">${r.start_mv??(r.direction==="auf"?t("matrix.open"):"–")} ${r.direction==="auf"?"↑ "+(r.mv_top??"–"):"↓ "+(r.mv_floor??"–")} mV</td>
     <td>${r.vmin!=null?`<b>${r.vmin} mV</b>`:"–"}</td><td>${F.ths(r.ths)}</td><td>${F.jth(r.jth)}</td><td class="wrap">${info}</td></tr>`;
    if(open)for(const x of r.sub){const sc=SUBCLS[x.status]||"bad";
      const what=x.status==="aktiv"?`${phase(x.phase)} · ${F.int(x.elapsed_s)} s`:
        x.status==="vorab_ok"?t("matrix.sub_precheck_ok",{p:x.frac?F.int(x.frac*100):"–"}):
        x.status==="gueltig"?t("matrix.sub_valid",{p:x.frac?F.int(x.frac*100):"–"})+(x.vmin?" · "+t("legend.vmin"):""):`<span class=clip>${esc(trMsg(x.reason)||t("status.rejected"))}</span>`;
      h+=`<tr class="sub s-${sc}${x.vmin?" vmin":""}"><td>${x.mv} mV${x.tag?` <span class=muted>(${esc(TAGKEY[x.tag]?t(TAGKEY[x.tag]):x.tag)})</span>`:""}</td><td class="muted">${x.stage?t("matrix.stage",{n:x.stage}):""}</td><td></td><td></td>
       <td>${F.ths(x.ths)}</td><td>${F.jth(x.jth)}</td><td class="wrap">${what}</td></tr>`}}
  $("mxtab").innerHTML=h+"</tbody>";
  document.querySelectorAll("#mxtab .toggle").forEach(e=>e.onclick=()=>{mxOpen[e.dataset.f]=!mxOpen[e.dataset.f];renderMatrix(M)});}

/* ================= Live / Zustand ================= */
const WSRC={estimated:"wallsrc.estimated",measured:"wallsrc.measured",mixed:"wallsrc.mixed",legacy:"wallsrc.legacy",
  geschaetzt:"wallsrc.estimated",gemessen:"wallsrc.measured",gemischt:"wallsrc.mixed"};   /* deutsch: Altlaeufe */
const isMeas=q=>q==="measured"||q==="gemessen";
const wsrc=q=>q?(WSRC[q]?t(WSRC[q]):q):"";
function setSrc(q){const e=$("t_wsrc");e.textContent=wsrc(q);e.className="tag"+(isMeas(q)?" meas":q?" est":"")}
function bestIdx(res){let b=null;for(const r of res){if(r.valid!==true||r.jth_local==null)continue;if(!b||r.jth_local<b.jth_local)b=r}return b?b.idx:null}
const RUNST={starting:"runstatus.ready",running:"runstatus.running",finished:"runstatus.finished",stopped:"runstatus.stopped",emergency_stop:"runstatus.emergency",deadman:"runstatus.deadman"};
const runst=s=>s?(RUNST[s]?t(RUNST[s]):(String(s).startsWith("error")?t("runstatus.error"):s)):t("runstatus.ready");
let last=null,pinIsSet=false;
function render(s){last=s;const C=s.control,S=s.sweep||{},cur=S.current||{},L=cur.live||{},A=cur.avg||{};
  $("miner").textContent=hostOf(C.miner_url);$("sstat").textContent=C.sweep_running?t("runstatus.running"):runst(S.status);
  if(!$("url").value)$("url").value=hostOf(C.miner_url);
  $("pinhint").textContent=C.pin_default?t("hint.pin_not_set_short"):"";
  if(C.pin_default!==!pinIsSet)refreshAuth();
  $("conn_sub").textContent=(C.connected_at?t("conn.connected_since",{t:C.connected_at.slice(11,16)}):minerSeen?t("conn.status_read",{t:minerSeen.toTimeString().slice(0,5)}):t("conn.not_connected"))+" · "+(C.pin_default?t("conn.pin_not_set"):t("conn.pin_active"));
  $("ebanner").innerHTML=C.emergency?`<div class="note bad" style="margin-bottom:18px;font-weight:700">${t("hint.estop_active",{t:esc(C.emergency_at||"")})}</div>`:"";
  $("start").disabled=!lastPlan||!pinIsSet||C.sweep_running||C.emergency;$("stop").disabled=!C.sweep_running;
  if(C.profile&&!minerData)fillMiner(C.profile,C.miner_url);
  const ri=C.run_info;$("live_sub").textContent=C.sweep_running?[cur.idx!=null?t("live.point",{n:cur.idx+1}):null,ri&&t("mode."+ri.sweep_mode),ri&&t(ri.full_measure?"live.full":"live.fast")].filter(Boolean).join(" · "):t("live.idle_sub");
  if(C.sweep_running){$("t_f").textContent=cur.frequency_mhz??"–";$("t_mv").textContent=cur.core_mv??"–";
    $("t_h").textContent=F.ths(L.local_ths);$("t_w").textContent=F.w(L.wall_w);setSrc(L.wall_src||s.wall_power_quelle);
    $("t_jk").textContent=t("tile.jth_window");$("t_j").textContent=F.jth(A.jth_local);$("t_e").textContent=F.pct(A.error_pct);
    $("t_t").innerHTML=F.temp(L.asic_c)+"<span class=u>&thinsp;/&thinsp;</span>"+F.temp(L.vr_c);$("t_v").textContent=F.volt(L.vin);
    if(minerData&&cur.frequency_mhz)$("f_now").innerHTML=`${cur.frequency_mhz} <small>MHz</small> / ${cur.core_mv} <small>mV</small>`;}
  $("t_ph").textContent=phase(cur.phase);const tot=(cur.elapsed_s||0)+(cur.remaining_s||0);
  $("t_pt").textContent=cur.idx!=null&&C.sweep_running?t("live.progress",{a:F.int(cur.elapsed_s),b:F.int(cur.remaining_s),n:A.samples??0}):"";
  $("bar").style.width=(tot>0&&C.sweep_running?100*cur.elapsed_s/tot:0)+"%";
  const res=S.results||[],best=bestIdx(res);
  $("restab").innerHTML=`<thead><tr><th>${t("col.point")}</th><th>MHz</th><th>mV</th><th>${t("col.meas_mv")}</th><th>TH/s</th><th>${t("col.pct_exp")}</th><th>${t("col.wall")}</th><th>${t("col.source")}</th><th>J/TH</th><th>${t("col.err")}</th><th>${t("col.temp_max")}</th><th class=wrap>${t("col.status")}</th></tr></thead><tbody>`+
   (res.map(r=>`<tr class="${r.idx===best?"s-ok":(r.valid===false?"s-bad":"")}"><td>${esc(r.label)}</td><td>${r.frequency_mhz}</td><td>${r.core_mv}</td><td>${F.temp(r.measured_core_mv)}</td>
   <td>${F.ths(r.hashrate_local_ths)}</td><td>${r.frac_of_expected?F.int(r.frac_of_expected*100):"–"}</td><td>${F.w(r.wall_avg)}</td><td>${esc(wsrc(r.wall_power_source))}</td>
   <td>${F.jth(r.jth_local)}</td><td>${F.pct(r.error_pct)}</td><td>${F.temp(r.asic_temp_max)}/${F.temp(r.vr_temp_max)}</td>
   <td class="wrap">${r.valid?t("word.valid"):t("word.invalid")}${r.idx===best?" · "+t("matrix.best_jth"):""} <span class=muted>${esc(trMsg(r.reason))}</span></td></tr>`).join("")||`<tr><td class=muted colspan=12>${t("matrix.no_points")}</td></tr>`)+"</tbody>";
  renderMatrix(s.matrix);
  renderLog(S.log||[],(C.events||[]).map(e=>"[ctrl] "+e));
  chartData=res;chartBest=best;redrawChart();}

/* ================= Log: Frontend-Puffer, Kopieren, Leeren ================= */
/* Der Server liefert nur ein gleitendes Fenster (Sweep-Log + Steuer-Ereignisse). Das Frontend sammelt neue
   Zeilen per Ueberlappungsvergleich in einem eigenen Puffer (max. LOG_KEEP Zeilen je Quelle).
   "Leeren" leert nur diesen Puffer/die Anzeige - Server-Log und journalctl bleiben unberuehrt. */
const LOG_KEEP=500;
const logAcc={sweep:{prev:[],lines:[]},ctrl:{prev:[],lines:[]}};
function logMerge(a,cur){const prev=a.prev;let k=Math.min(prev.length,cur.length);
  for(;k>0;k--){let ok=true;for(let i=0;i<k;i++)if(prev[prev.length-k+i]!==cur[i]){ok=false;break}if(ok)break}
  a.lines=a.lines.concat(cur.slice(k)).slice(-LOG_KEEP);a.prev=cur.slice();}
function logText(){return [...logAcc.sweep.lines,...logAcc.ctrl.lines].join("\n")}
function renderLog(sweep,ctrl){logMerge(logAcc.sweep,sweep);logMerge(logAcc.ctrl,ctrl);
  const el=$("log"),txt=logText();if(el.textContent!==txt)el.textContent=txt;}
function copyFallback(txt){const ta=document.createElement("textarea");ta.value=txt;ta.setAttribute("readonly","");
  ta.style.position="fixed";ta.style.top="-1000px";ta.style.opacity="0";document.body.appendChild(ta);ta.select();
  let ok=false;try{ok=document.execCommand("copy")}catch(e){ok=false}document.body.removeChild(ta);return ok}
async function copyLog(){const txt=logText(),b=$("logcopy");let ok=false;
  if(navigator.clipboard&&window.isSecureContext){try{await navigator.clipboard.writeText(txt);ok=true}catch(e){ok=false}}
  if(!ok)ok=copyFallback(txt);
  if(ok){b.textContent=t("log.copied");clearTimeout(b._t);b._t=setTimeout(()=>b.textContent=t("log.copy"),1500)}
  else toast(t("log.copy")+": ✗",true);}
function clearLog(){for(const a of Object.values(logAcc))a.lines=[];$("log").textContent="";}
$("logcopy").onclick=copyLog;$("logclear").onclick=clearLog;

/* ================= J/TH-Kurve ================= */
let chartData=[],chartBest=null;
function redrawChart(){drawChart(chartData,chartBest)}
function drawChart(results,best){const cv=$("chart");if(!cv)return;const dpr=window.devicePixelRatio||1,W=cv.clientWidth,H=cv.clientHeight;if(!W)return;
  cv.width=W*dpr;cv.height=H*dpr;const g=cv.getContext("2d");g.scale(dpr,dpr);g.clearRect(0,0,W,H);
  const cs=getComputedStyle(root),muted=cs.getPropertyValue("--muted").trim(),line=cs.getPropertyValue("--line").trim(),acc=cs.getPropertyValue("--accent").trim(),text=cs.getPropertyValue("--text").trim();
  const pts=(results||[]).filter(r=>r.valid===true&&r.jth_local!=null);g.font="13px system-ui,sans-serif";g.fillStyle=muted;
  if(!pts.length){g.fillText(t("curve.empty"),18,28);$("curveLegend").innerHTML="";return}
  let x0=Math.min(...pts.map(r=>r.core_mv)),x1=Math.max(...pts.map(r=>r.core_mv)),y0=Math.min(...pts.map(r=>r.jth_local)),y1=Math.max(...pts.map(r=>r.jth_local));
  if(x1-x0<40){const m=(x0+x1)/2;x0=m-20;x1=m+20}if(y1-y0<0.5){const m=(y0+y1)/2;y0=m-0.25;y1=m+0.25}const py=(y1-y0)*.12;y0-=py;y1+=py;const px=(x1-x0)*.05;x0-=px;x1+=px;
  const Lm=48,R=14,T=12,B=34,X=v=>Lm+(v-x0)/(x1-x0)*(W-Lm-R),Y=v=>T+(1-(v-y0)/(y1-y0))*(H-T-B);
  g.font="12px system-ui,sans-serif";g.strokeStyle=line;g.lineWidth=1;g.textAlign="right";g.textBaseline="middle";
  for(let i=0;i<=5;i++){const v=y0+(y1-y0)*i/5;g.beginPath();g.moveTo(Lm,Y(v));g.lineTo(W-R,Y(v));g.stroke();g.fillText(fmt(v,2),Lm-6,Y(v))}
  g.textAlign="center";g.textBaseline="top";for(let i=0;i<=5;i++){const v=x0+(x1-x0)*i/5;g.fillText(Math.round(v),X(v),H-B+6)}
  g.fillText("core_mv",(Lm+W-R)/2,H-15);
  const fs=[...new Set(pts.map(r=>r.frequency_mhz))].sort((a,b)=>a-b),leg=[];
  fs.forEach((f,i)=>{const hue=Math.round(205-175*(fs.length>1?i/(fs.length-1):0)),col=`hsl(${hue},72%,55%)`;leg.push([f,col]);
    const s=pts.filter(r=>r.frequency_mhz===f).sort((a,b)=>a.core_mv-b.core_mv);g.strokeStyle=col;g.lineWidth=1.6;g.beginPath();
    s.forEach((r,j)=>j?g.lineTo(X(r.core_mv),Y(r.jth_local)):g.moveTo(X(r.core_mv),Y(r.jth_local)));g.stroke();
    s.forEach(r=>{g.fillStyle=col;g.beginPath();g.arc(X(r.core_mv),Y(r.jth_local),3.4,0,7);g.fill()})});
  const b=pts.find(r=>r.idx===best);
  if(b){const bx=X(b.core_mv),by=Y(b.jth_local);g.strokeStyle=acc;g.lineWidth=2.5;g.beginPath();g.arc(bx,by,9,0,7);g.stroke();
    g.fillStyle=text;g.textAlign="left";g.textBaseline="bottom";g.font="600 12px system-ui,sans-serif";
    g.fillText(t("curve.best",{f:b.frequency_mhz,mv:b.core_mv,j:F.jth(b.jth_local)}),Math.min(bx+12,W-260),by-8)}
  $("curveLegend").innerHTML=leg.map(([f,c])=>`<span><i style="display:inline-block;width:10px;height:10px;border-radius:50%;background:${c};margin-right:5px"></i>${f} MHz</span>`).join("");}
window.addEventListener("resize",redrawChart);

/* ================= Historie ================= */
const SSKEY={efficiency:"ss.efficiency",knee:"ss.knee",compromise:"ss.compromise",target:"ss.target"};
let histBusy=false;
async function loadHist(){if(histBusy)return;histBusy=true;try{const j=await api("GET","/runs");$("hist_sub").textContent=t("hist.count",{n:j.runs.length});
  $("hist").innerHTML=j.runs.map(r=>{const S=r.sweet_spots||{};const short=String(r.run_tag||r.id).includes("kurz");
    const badge=short?`<span class="badge short">${t("hist.badge_short")}</span>`:r.status==="finished"?`<span class="badge fin">${t("hist.badge_done")}</span>`:`<span class="badge other">${esc(runst(r.status))}</span>`;
    const ss=Object.keys(SSKEY).filter(k=>S[k]).map(k=>`<div class="subcard${k==="efficiency"?" best":""}"><div class="k">${t(SSKEY[k])}</div><div class="v">${S[k].freq} MHz / ${S[k].mv} mV</div>
      <div class="d">${F.ths(S[k].hashrate_local_ths)} TH/s · ${F.w(S[k].wall_avg)} W · ${F.jth(S[k].jth_local)} J/TH</div></div>`).join("");
    const rid=esc(r.id);
    return `<div class="subcard"><div class="run-h"><div><div class="title">${rid} ${badge}</div>
      <div class="muted" style="font-size:12.5px">${esc(r.timestamp||"")} → ${esc(r.finished_at||"")}${r.sweep_mode?" · "+t("mode."+r.sweep_mode):""} · ${t("hist.valid_of",{a:r.valid_points,b:r.points})}${r.profile?" · "+esc(r.profile):""}</div></div>
      <div class="row">${["csv","json","xlsx","pdf"].map(x=>`<a class="btn ghost small" href="/export/${encodeURIComponent(r.id)}.${x}">${x.toUpperCase()}</a>`).join("")}</div></div>
      ${short?`<div class="note warn" style="margin-bottom:10px" title="${esc(r.note||"")}">${t("hist.short_note")}</div>`:(r.note?`<div class="note warn" style="margin-bottom:10px">${esc(r.note)}</div>`:"")}
      ${ss?`<div class="ss">${ss}</div>`:`<div class="muted">${t("hist.no_ss")}</div>`}
      <div class="bestfor"><div class="row">
        <label class="field"><span>${t("bf.label")}</span><input type="text" inputmode="decimal" class="bf-in" data-run="${rid}" placeholder="${t("bf.placeholder")}" style="width:110px"></label>
        <button type="button" class="btn ghost bf-btn" data-run="${rid}">${t("btn.best_for")}</button></div>
        <div class="bf-out" data-run="${rid}"></div></div></div>`}).join("")||`<span class=muted>${t("hist.none")}</span>`;
  document.querySelectorAll(".bf-btn").forEach(b=>b.onclick=()=>bestFor(b.dataset.run));
  document.querySelectorAll(".bf-in").forEach(i=>i.onkeydown=e=>{if(e.key==="Enter")bestFor(i.dataset.run)});}
  catch(e){$("hist").textContent=trErr(e)}finally{histBusy=false}}

/* ================= Ziel-Hashrate: besten Punkt finden (lesend) + uebernehmen (PIN) ================= */
const ROLEKEY={reserve_recommended:"bf.role.reserve_recommended",vmin_no_reserve:"bf.role.vmin_no_reserve",vmin_alternative:"bf.role.vmin_alternative",reserve:"bf.role.reserve",vmin:"bf.role.vmin"};
function bfSel(cls,run){return [...document.querySelectorAll("."+cls)].find(e=>e.dataset.run===run)}
function srcTag(q){return `<span class="tag ${isMeas(q)?"meas":"est"}">${esc(wsrc(q))}</span>`}
function ptLine(p){return `${p.freq} MHz / ${p.mv} mV · ${F.ths(p.hashrate_ths)} TH/s · ${F.w(p.wall_w)} W ${srcTag(p.wall_power_quelle)} · ${F.jth(p.jth_local)} J/TH`}
async function bestFor(run){const inp=bfSel("bf-in",run),out=bfSel("bf-out",run);const tv=Number(String(inp.value).replace(",","."));
  if(!(tv>0)){out.innerHTML=`<div class="note bad">${t("bf.enter_target")}</div>`;return}
  out.innerHTML=`<span class="muted">${t("bf.searching")}</span>`;
  try{const r=await api("POST","/best-for",{run_id:run,target_ths:tv},false);let h="";
    if(r.caution_code||r.caution)h+=`<div class="note warn">${t("bf.caution."+(r.caution_code||"short_run"))}</div>`;
    if(r.chosen){const c=r.chosen;
      h+=`<div class="bf-chosen"><div class="k">${t("bf.best_for",{t:F.num(tv)})} · ${t(ROLEKEY[c.role]||"bf.role.vmin")}</div>
        <div class="bf-big">${c.freq} <small>MHz</small> / ${c.mv} <small>mV</small></div>
        <div class="bf-vals"><span><b>${F.ths(c.hashrate_ths)}</b> TH/s</span><span><b>${F.jth(c.jth_local)}</b> J/TH</span><span><b>${F.w(c.wall_w)}</b> W ${srcTag(c.wall_power_quelle)}</span></div>
        <div class="row needpin" style="${pinIsSet?"":"display:none"};margin-top:10px">
          <button type="button" class="btn" data-apply="live">${t("btn.apply_live")}</button>
          <button type="button" class="btn ghost" data-apply="save">${t("btn.apply_save")}</button></div>
        <div class="bf-apply muted" style="margin-top:6px"></div></div>`;
      if(r.vmin_alternative)h+=`<div class="bf-small">${t("bf.vmin_alt")}: ${ptLine(r.vmin_alternative)}</div>`;}
    else h+=`<div class="note bad">${r.message_code?t("bf.msg."+r.message_code,{t:F.num((r.message_params||{}).target),top:F.ths((r.message_params||{}).top)}):esc(r.message||"")}</div>`;
    if(r.below)h+=`<div class="bf-small">${t("bf.below")}: ${ptLine(r.below)}</div>`;
    out.innerHTML=h;
    out.querySelectorAll("[data-apply]").forEach(b=>b.onclick=()=>applyPoint(r.chosen,b.dataset.apply==="save",out.querySelector(".bf-apply")));
  }catch(e){out.innerHTML=`<div class="note bad">${esc(trErr(e))}</div>`}}
async function applyPoint(c,save,box){
  if(save&&!confirm(t("confirm.save",{f:c.freq,mv:c.mv})))return;
  const body={frequency_mhz:c.freq,core_mv:c.mv,save:!!save};if(save)body.confirm_save=true;
  box.textContent=t(save?"bf.saving":"bf.setting");
  try{const r=await api("POST","/point/set",body,true);
    box.innerHTML=`<b>${t(save?"bf.saved":"bf.set_live")}</b> · ${t("bf.verify")}: ${r.verify_ok?`<b style='color:var(--st-ok-fg)'>${t("bf.verify_ok")}</b>`:`<b style='color:var(--st-bad-fg)'>${t("bf.verify_fail")}</b>`}
      · ${t("bf.miner_reports",{f:r.current_frequency_mhz,mv:F.int(r.measured_core_mv)})}`;
    toast(t(save?"msg.saved_point":"msg.set_point",{f:c.freq,mv:c.mv}),!r.verify_ok);minerData=null;pollIdle(true);
  }catch(e){box.innerHTML=`<b style="color:var(--st-bad-fg)">${esc(trErr(e))}</b>`;toast(trErr(e),true)}}

/* ================= Ruhezustand: Miner-Status lesen (nur GET) ================= */
let lastIdle=null;
function showIdle(j){const st=j.status;fillMiner(st,j.miner_url);
  $("t_f").textContent=st.current_frequency_mhz??"–";$("t_mv").textContent=st.core_mv??"–";$("t_h").textContent=F.ths(st.local_hashrate_ghs/1000);
  $("t_w").textContent=F.w(j.wall_power_w);setSrc(j.wall_power_quelle);$("t_jk").textContent=t("tile.jth_now");
  $("t_j").textContent=(j.wall_power_w&&st.local_hashrate_ghs)?F.jth(j.wall_power_w/(st.local_hashrate_ghs/1000)):"–";
  $("t_e").textContent="–";$("t_t").innerHTML=F.temp(st.asic_temp_c)+"<span class=u>&thinsp;/&thinsp;</span>"+F.temp(st.vr_temp_c);$("t_v").textContent=F.volt(st.input_voltage_v)}
async function pollIdle(once){if(!last||!last.control.sweep_running){try{const j=await api("GET","/status.json");lastIdle=j;showIdle(j)}
  catch(e){lastIdle=null;setMinerOff(t("miner.unreachable",{why:esc(trErr(e))}))}}
  if(once!==true)setTimeout(pollIdle,10000)}
let wasRunning=false;
async function poll(){try{const s=await api("GET","/state.json");render(s);if(wasRunning&&!s.control.sweep_running)loadHist();wasRunning=s.control.sweep_running}catch(e){$("sstat").textContent=t("msg.no_server")}setTimeout(poll,1000)}

/* ================= PIN einrichten / aendern ================= */
async function refreshAuth(){try{const j=await api("GET","/auth-status");pinIsSet=!!j.pin_set;
  $("pinsetup").style.display=pinIsSet?"none":"";$("pinchange").style.display=pinIsSet?"":"none";
  document.querySelectorAll(".needpin").forEach(e=>e.style.display=pinIsSet?"":"none")}catch(e){}}
function pinMsg(id,text,ok){const e=$(id);e.className="pinmsg "+(ok?"ok":"bad");e.textContent=text}
function checkNewPin(a,b){if(a!==b)return t("msg.pins_mismatch");if(a.length<4)return t("msg.pin_too_short");
  if(a==="CHANGEME")return t("msg.pin_changeme");return null}
function adoptPin(p){$("pin").value=p;lsSet("ctrl_pin",p);$("pinerr").textContent=""}
document.querySelectorAll(".showpin").forEach(cb=>cb.onchange=()=>{document.querySelectorAll(`[id^="${cb.dataset.scope}_"].pinf`).forEach(i=>i.type=cb.checked?"text":"password")});
$("ps_btn").onclick=async()=>{const a=$("ps_new").value.trim(),b=$("ps_rep").value.trim(),err=checkNewPin(a,b);
  if(err){pinMsg("ps_msg",err,false);return}
  try{await api("POST","/set-pin",{new_pin:a},false);adoptPin(a);$("ps_new").value=$("ps_rep").value="";
    pinMsg("ps_msg",t("msg.pin_set"),true);toast(t("msg.pin_set"));await refreshAuth()}catch(e){pinMsg("ps_msg",trErr(e),false)}};
$("pc_btn").onclick=async()=>{const c=$("pc_cur").value.trim(),a=$("pc_new").value.trim(),b=$("pc_rep").value.trim(),err=checkNewPin(a,b);
  if(!c){pinMsg("pc_msg",t("msg.enter_current_pin"),false);return}
  if(err){pinMsg("pc_msg",err,false);return}
  try{await api("POST","/set-pin",{current_pin:c,new_pin:a},false);adoptPin(a);$("pc_cur").value=$("pc_new").value=$("pc_rep").value="";
    pinMsg("pc_msg",t("msg.pin_set"),true);toast(t("msg.pin_set"));await refreshAuth()}
  catch(e){pinMsg("pc_msg",e.status===403?t("msg.current_pin_wrong"):trErr(e),false)}};

/* ================= Start ================= */
applyTheme(lsGet("ctrl_theme")==="light"?"light":"dark");
applyAccent(["blue","green","violet","orange"].includes(lsGet("ctrl_accent"))?lsGet("ctrl_accent"):"blue");
initLang().then(()=>{refreshAuth();poll();pollIdle()});
</script>
</body>
</html>
"""


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Miner Steuerzentrale")
    ap.add_argument("--config", default=os.path.join(HERE, "config.json"))
    ap.add_argument("--reset-pin", action="store_true",
                    help="Notausgang: control_pin in config.json atomar auf CHANGEME zuruecksetzen und beenden "
                         "(laufenden Server danach mit 'kill -HUP <pid>' neu laden lassen)")
    args = ap.parse_args()
    if args.reset_pin:
        write_pin_to_config(DEFAULT_PIN, os.path.abspath(args.config), update_runtime=False)
        print(f"control_pin in {os.path.abspath(args.config)} reset to initial setup (CHANGEME).")
        return

    CONFIG_PATH["path"] = os.path.abspath(args.config)
    cfg, created = load_config(args.config)
    CFG.update(cfg)
    apply_config()
    print(f"Config {'CREATED (defaults)' if created else 'loaded'}: {args.config}", flush=True)
    print(f"  miner_url={CFG['miner_url']}  band freq_low_pct={CFG['freq_low_pct']} floor_pct={CFG['floor_pct']}  "
          f"window {CFG['window_min_s']}-{CFG['window_max_s']}s/{CFG['target_hits']} hits  "
          f"HARD {CFG['hard_asic_c']}/{CFG['hard_vr_c']} C", flush=True)
    if not pin_set():
        print("  PIN: NOT SET -> write actions (except emergency stop) locked. "
              "Set it in the web UI under 'Set control PIN' (no restart needed).", flush=True)
    else:
        print("  PIN: set - write actions require header X-Control-Pin", flush=True)
    print(f"  Idle: no write access to the miner (GET only for /status.json, /connect, /plan) | UI default language: {CFG.get('language')}", flush=True)

    host, port = CFG["bind_host"], int(CFG["bind_port"])
    if not port_free(host, port):
        print(f"ERROR: {host}:{port} is in use - stop the other process or change bind_port.", flush=True)
        sys.exit(1)
    server = ThreadingHTTPServer((host, port), Handler)
    server.daemon_threads = True
    threading.Thread(target=server.serve_forever, name="http", daemon=True).start()
    print(f"Control center running: http://{host}:{port}", flush=True)
    ip = lan_ip()
    if ip:
        print(f"Reachable on LAN: http://{ip}:{port}", flush=True)

    stop = threading.Event()

    def _term(signum, frame):
        stop.set()
    signal.signal(signal.SIGTERM, _term)
    signal.signal(signal.SIGHUP, lambda signum, frame: reload_pin_from_file())
    try:
        while not stop.is_set():
            stop.wait(1)
    except KeyboardInterrupt:
        pass
    finally:
        if sweep_running():
            print("Prozessende: laufenden Sweep stoppen und auf Ausgangspunkt zuruecksetzen ...", flush=True)
            for c in contexts.values():
                if c.sweep_running():
                    c.stop_event.set()
                    c.thread.join(timeout=120)
        server.shutdown()
        server.server_close()
        print("Beendet.", flush=True)


if __name__ == "__main__":
    main()
